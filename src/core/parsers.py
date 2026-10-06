from __future__ import annotations

import base64
import math
import mimetypes
import os
import re
from io import BytesIO
from pathlib import Path
from typing import Any

from .models import ExtractedSource


SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".pptx", ".png", ".jpg", ".jpeg", ".txt"}
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg"}


def extract_source(
    filename: str, content: bytes, pasted_text: str | None = None
) -> ExtractedSource:
    """Extract text from a local source while returning visible status for every failure mode."""

    filename = filename or "pasted.txt"
    suffix = Path(filename).suffix.lower()
    mime_type = mimetypes.guess_type(filename)[0] or "application/octet-stream"

    if pasted_text is not None:
        text = pasted_text.strip()
        return ExtractedSource(
            filename=filename,
            mime_type="text/plain",
            size_bytes=len(pasted_text.encode("utf-8")),
            text=text,
            status="ok" if text else "empty",
            error_message=None if text else "粘贴文本为空。",
        )

    if suffix not in SUPPORTED_EXTENSIONS:
        return ExtractedSource(
            filename=filename,
            mime_type=mime_type,
            size_bytes=len(content or b""),
            text="",
            status="unsupported",
            error_message=f"暂不支持 {suffix or '未知'} 文件类型。",
        )

    try:
        if suffix == ".txt":
            return _txt(filename, mime_type, content)
        if suffix == ".pdf":
            return _pdf(filename, mime_type, content)
        if suffix == ".docx":
            return _docx(filename, mime_type, content)
        if suffix == ".pptx":
            return _pptx(filename, mime_type, content)
        if suffix in IMAGE_EXTENSIONS:
            return _image(filename, mime_type, content)
    except Exception as exc:
        return ExtractedSource(
            filename=filename,
            mime_type=mime_type,
            size_bytes=len(content or b""),
            text="",
            status="parse_error",
            error_message=f"解析失败：{exc}",
        )
    raise AssertionError("unreachable")


def _txt(filename: str, mime_type: str, content: bytes) -> ExtractedSource:
    for encoding in ("utf-8", "utf-8-sig", "gb18030"):
        try:
            text = content.decode(encoding).strip()
            break
        except UnicodeDecodeError:
            continue
    else:
        text = content.decode("utf-8", errors="ignore").strip()
    return ExtractedSource(
        filename=filename,
        mime_type=mime_type,
        size_bytes=len(content),
        text=text,
        status="ok" if text else "empty",
        error_message=None if text else "文本文件为空。",
    )


def _is_scanned_page(page) -> bool:
    """检测页面是否为纯扫描页（无文本层）"""
    text = page.get_text("text").strip()
    if len(text) > 10:
        return False
    images = page.get_images()
    drawings = page.get_drawings()
    return bool(images or drawings)


_tesseract_checked = False
_tesseract_available_flag = False


def _tesseract_available() -> bool:
    """一次性检测 Tesseract 是否真正可用（避免未安装时每页白跑 0.5s 异常）。"""
    global _tesseract_checked, _tesseract_available_flag
    if _tesseract_checked:
        return _tesseract_available_flag
    _tesseract_checked = True
    try:
        import pytesseract
    except ImportError:
        return False
    try:
        from .tesseract_auto import setup_pytesseract

        setup_pytesseract()
    except Exception:
        pass
    try:
        pytesseract.get_tesseract_version()
        _tesseract_available_flag = True
    except Exception:
        _tesseract_available_flag = False
    return _tesseract_available_flag


def _ocr_page_pymupdf(page, dpi: int = 300) -> str:
    """用 PyMuPDF + pytesseract OCR 识别单页（自动配置tesseract路径）"""
    if not _tesseract_available():
        return ""

    try:
        import pytesseract
    except ImportError:
        return ""

    try:
        from PIL import Image
    except ImportError:
        return ""

    # auto-configure tesseract path on first call
    if not getattr(_ocr_page_pymupdf, "_tess_configured", False):
        try:
            from .tesseract_auto import setup_pytesseract

            setup_pytesseract()
        except Exception:
            pass
        _ocr_page_pymupdf._tess_configured = True

    try:
        import pymupdf as fitz
    except Exception:
        import fitz

    mat = fitz.Matrix(dpi / 72, dpi / 72)
    pix = page.get_pixmap(matrix=mat)
    img = Image.open(BytesIO(pix.tobytes("png")))
    try:
        text = pytesseract.image_to_string(img, lang="chi_sim+eng")
        return text.strip()
    except Exception:
        return ""


def _ocr_page_paddle(page, dpi: int = 300) -> str:
    """用 PaddleOCR 识别单页（备用）"""
    try:
        from paddleocr import PaddleOCR
    except ImportError:
        return ""

    try:
        import pymupdf as fitz
    except Exception:
        import fitz

    mat = fitz.Matrix(dpi / 72, dpi / 72)
    pix = page.get_pixmap(matrix=mat)
    img_bytes = pix.tobytes("png")

    tmp_path = f"_tmp_ocr_page_{page.number}.png"
    try:
        with open(tmp_path, "wb") as f:
            f.write(img_bytes)

        ocr = PaddleOCR(use_angle_cls=True, lang="ch", show_log=False)
        result = ocr.ocr(tmp_path, cls=True)
    finally:
        if os.path.exists(tmp_path):
            os.remove(tmp_path)

    if not result or not result[0]:
        return ""
    lines = [line[1][0] for line in result[0]]
    return "\n".join(lines)


def _pdf(filename: str, mime_type: str, content: bytes) -> ExtractedSource:
    try:
        import pymupdf as fitz
    except Exception:
        try:
            import fitz
        except Exception as exc:
            return ExtractedSource(
                filename=filename,
                mime_type=mime_type,
                size_bytes=len(content),
                text="",
                status="parse_error",
                error_message=f"Missing PyMuPDF; cannot parse PDF: {exc}",
            )

    try:
        doc = fitz.open(stream=content, filetype="pdf")
    except Exception as exc:
        msg = str(exc).lower()
        if "password" in msg or "encrypt" in msg:
            hint = "PDF 已加密或受密码保护，请先解除保护后重试。"
        elif "corrupt" in msg or "invalid" in msg or "not a pdf" in msg:
            hint = "PDF 文件已损坏或格式无效，请重新下载或转换后重试。"
        elif "memory" in msg:
            hint = "PDF 文件过大导致内存不足，请尝试拆分为较小文件后重试。"
        else:
            hint = f"PDF 无法打开：{exc}"
        return ExtractedSource(
            filename=filename,
            mime_type=mime_type,
            size_bytes=len(content),
            text="",
            status="parse_error",
            error_message=hint,
        )

    try:
        with doc:
            page_texts = []
            figures: list[dict[str, Any]] = []
            math_region_formulas: list[dict[str, Any]] = []
            scan_count = 0
            fig_start_id = 0

            # 章节标题：PDF 大纲书签优先（level 归一化到「1=题名、2=节、3=小节」），字号启发式兜底。
            # 每条候选都带真实页码（最终写成 [[H:level@page]] 标记）——页码只有解析器知道，
            # 交给模型去猜只会得到一串编造值。
            outline_headings = _outline_headings(doc)
            headings = _normalize_heading_levels(
                [
                    (level, title, page)
                    for page, entries in outline_headings.items()
                    for (level, title) in entries
                ]
            )

            for page_index, page in enumerate(doc):
                raw = page.get_text("text")
                if raw.strip():
                    page_text = raw
                elif _is_scanned_page(page):
                    # 扫描页文本由上层（api_server 快速通道）RapidOCR 统一识别，
                    # 此处不再逐页跑 Tesseract/PaddleOCR：未装时白等异常，装了会拖到分钟级。
                    scan_count += 1
                    page_text = ""
                else:
                    page_text = raw

                if raw.strip():
                    headings.extend(_font_heading_lines(page, page_index + 1))

                # 抽全部位图并按 y 坐标原位插入 [[FIG:id]] 标记（无标题也保留）
                # dropped_figs：本页被判为伪图（框住正文 / 水印层）而剔除的候选。
                # 只要剔过东西，就说明这页的候选已经找过了，后面不能再触发按图注兜底
                # —— 否则「剔除伪图」会反过来凭空多出一批兜底框取的新图。
                dropped_figs: list[dict[str, Any]] = []
                page_figs = _page_figures(
                    doc, page, page_index, fig_start_id, dropped=dropped_figs
                )
                fig_start_id += len(page_figs)
                # 矢量图检测：在已有位图之外寻找矢量绘制区域并渲染
                bitmap_bboxes = [f["bbox"] for f in page_figs if f.get("bbox")]
                vector_figs = _vector_figure_regions(
                    page,
                    page_index,
                    fig_start_id,
                    bitmap_bboxes,
                    max_figures=_MAX_FIGURES - len(page_figs),
                    dropped=dropped_figs,
                )
                fig_start_id += len(vector_figs)
                page_figs = page_figs + vector_figs

                # 若页面存在"图 N"标题但尚未提取到位图或矢量图，执行标题驱动截取兜底
                if not page_figs and not dropped_figs:
                    caption_figs = _caption_driven_figure_fallback(
                        doc, page, page_index, fig_start_id
                    )
                    fig_start_id += len(caption_figs)
                    page_figs.extend(caption_figs)

                # 整幅区域：按版面把同页同一幅多子图合并为整幅（含子图标签），挂 whole_ref
                for reg in _layout_figure_regions(page, page_figs):
                    _attach_whole_ref(page_figs, [reg])

                figures.extend(page_figs)
                if page_text.strip() and page_figs:
                    page_text = _insert_figure_markers(page_text, page, page_figs)

                # 多行碎片组成的显示公式：文本层线性化已毁结构，
                # 直接把整块区域渲染成图片，正文只留 [[MATH:id]] 占位符
                if raw.strip():
                    rendered_regions: list[
                        tuple[dict[str, Any], str, str, dict[str, Any]]
                    ] = []
                    for k, region in enumerate(
                        _display_math_regions(
                            page,
                            figure_bboxes=[
                                f["bbox"] for f in page_figs if f.get("bbox")
                            ],
                        )
                    ):
                        # clip 已含 2pt 白边并削去邻近正文/页眉残影
                        img = _render_figure_region(
                            page, region.get("clip") or region["bbox"], pad=0.0
                        )
                        if not img:
                            continue
                        placeholder = f"region-{page_index}-{k}"
                        rendered_regions.append(
                            (region, placeholder, f"[[MATH:{placeholder}]]", img)
                        )
                    # 标记按几何阅读顺序锚定：多重集替换的「首个命中行」锚定
                    # 会把共享碎片行（|||、|、z）的多个公式吸到同一位置
                    placed: set[str] = set()
                    if rendered_regions:
                        page_text, placed = _insert_math_region_markers(
                            page,
                            page_text,
                            [r for r, _, _, _ in rendered_regions],
                            [m for _, _, m, _ in rendered_regions],
                        )
                    for region, placeholder, marker, img in rendered_regions:
                        page_text, tail_nums, first_removed = _remove_math_region_lines(
                            page_text, region["lines"]
                        )
                        if marker not in placed:
                            if first_removed is None:
                                continue  # 碎片不在本页文本里，区域存疑，不建记录
                            lines = page_text.split("\n")
                            lines.insert(first_removed, marker)
                            page_text = "\n".join(lines)
                        # 编号已随区域渲染进图片：正文中独立成行的编号一并移除
                        region_nums = list(region.get("numbers") or [])
                        for num_txt in region_nums:
                            page_text = _replace_standalone_line(
                                page_text, num_txt, ""
                            )
                        # 行尾剥离下来的编号（未进图片）也作为记录编号，供前端角标
                        all_nums = [
                            re.search(r"\d{1,3}[a-zA-Z]?", t).group(0)
                            for t in (*region_nums, *tail_nums)
                            if re.search(r"\d{1,3}[a-zA-Z]?", t)
                        ]
                        num_match = all_nums[0] if all_nums else None
                        math_region_formulas.append(
                            {
                                "text": "\n".join(region["lines"]),
                                "raw": "\n".join(region["lines"]),
                                "formula": "",
                                "image": img["image"],
                                "width": img["width"],
                                "height": img["height"],
                                "source": "pymupdf-region-image",
                                "page": page_index + 1,
                                # number 置位以阻止编号挂载流程再挂别的编号；
                                # 前端凭 num_in_image 不渲染编号角标（编号已在图里）
                                "number": num_match,
                                "num_in_image": bool(region_nums),
                                "_pid": placeholder,
                            }
                        )

                page_texts.append(page_text)

            pymupdf_formulas = _formula_candidates_from_pymupdf(doc)
            # 注意：math_region_formulas 不并入 pymupdf_formulas——它们带整块区域截图，
            # 是最高质量记录，须排在去重输入首位（截断上限内优先保留）

        # 过滤期刊信息图、重排 fig id 并同步重写正文标记
        figures, page_texts = _finalize_figures_and_markers(figures, page_texts)
    except Exception as exc:
        return ExtractedSource(
            filename=filename,
            mime_type=mime_type,
            size_bytes=len(content),
            text="",
            status="parse_error",
            error_message=f"PDF 文本提取失败：{exc}",
        )

    # 页眉/页脚等处重复出现的行（期刊名、卷期、作者、DOI、页码等）直接剔除
    page_texts = _strip_repeated_header_footer(page_texts)
    # 独立成行的邮箱行（如首页页脚「*defu@bit.edu.cn」通信作者信箱）只出现一次，
    # 重复行过滤覆盖不到 → 单独剔除，避免被 reflow 拼进正文
    page_texts = [
        "\n".join(
            ln for ln in pt.split("\n") if not _EMAIL_LINE_RE.match(ln)
        )
        for pt in page_texts
    ]
    raw_text = "\n\n".join(pt.strip() for pt in page_texts if pt.strip()).strip()

    # 标题去重后打 [[H:n]] 标记（空行包裹，保证 reflow 后独立成段）
    text = _mark_headings_in_text(raw_text, _dedupe_headings(headings))

    formulas = _dedupe_formulas(
        [
            *math_region_formulas,
            *pymupdf_formulas,
            *_formula_candidates_from_pdfplumber(content),
            *_formula_candidates_from_text(raw_text),
        ]
    )
    # 给公式分配稳定 id，并在正文留 [[MATH:id]] 占位符，供前端用 KaTeX 原位渲染
    formulas = _assign_formula_ids(formulas)
    text = _mark_formulas_in_text(text, formulas)
    # 区域图片公式的占位符重写为分配后的稳定 id；去重被丢弃的记录残留占位符直接清除
    for rec in math_region_formulas:
        pid = rec.pop("_pid", None)
        if pid and rec.get("id"):
            text = text.replace(f"[[MATH:{pid}]]", f"[[MATH:{rec['id']}]]")
    text = re.sub(r"\[\[MATH:region-\d+-\d+\]\]", "", text)
    # 游离的公式编号（N）挂回公式条目（前端渲染到公式右侧）
    text = _attach_formula_numbers(text, formulas)
    # 清除多行公式残留的碎片行（须在 reflow 前，否则碎片会被拼进正文句子）
    text = _drop_formula_debris_lines(text)
    # 剔除误判为公式的正文片段（含 ≥3 个汉字），还原为正文
    text = _purge_bogus_formulas(text, formulas)
    # 把「一句还没写完就换行」的硬换行合并成自然段落
    text = _reflow_paragraphs(text).strip()
    # 清除公式/表格排版残留的竖线串（| | | 垃圾）
    text = _strip_pipe_artifacts(text)
    # 后处理补出来的标题标记（编号短标题兜底）没有页码，按标题行所在页回填
    text = _attach_heading_pages(text, page_texts)
    # 后处理（公式候选还原、碎片清行）可能把标题文本挪走，留下没有标题的
    # [[H:n]] 空标记 —— 阅读器会渲染成空标题，目录还会把下一行正文吞成标题
    text = _normalize_heading_markers(text)
    # 无图注的图挂上「所在章节标题」，供前端图表列表显示（不编造图号）
    attach_figure_contexts(text, figures)

    # 扫描件但未识别到文本：给出精准提示
    if not text and scan_count > 0:
        return ExtractedSource(
            filename=filename,
            mime_type=mime_type,
            size_bytes=len(content),
            text="",
            formulas=formulas,
            figures=figures,
            status="empty",
            ocr_used=True,
            error_message=f"PDF 为扫描件（共 {scan_count} 页），本地无文本层且未安装 Tesseract。系统会用视觉模型自动识别；若仍无结果，请在「API 设置 → 视觉模型」里配置 vision_model（如 glm-4v-flash）。",
        )

    return ExtractedSource(
        filename=filename,
        mime_type=mime_type,
        size_bytes=len(content),
        text=text,
        formulas=formulas,
        figures=figures,
        status="ok" if text or formulas or figures else "empty",
        ocr_used=scan_count > 0,
        error_message=None
        if text or formulas or figures
        else "PDF did not contain extractable text or formulas.",
    )


_SPACE_RE = re.compile(r"[\s\u00a0]+")
_DIGIT_RE = re.compile(r"\d")


def _strip_spaces(value: str) -> str:
    """去掉全部空白（含不换行空格）。"""
    return _SPACE_RE.sub("", value)


def _norm_line_from_strip(line_strip: str) -> str:
    """把 _strip_spaces 的结果补成 _norm_line 的归一形式（小写 + 连字符归一）。"""
    return line_strip.lower().replace("−", "-").replace("–", "-")


def _norm_line(line: str) -> str:
    """去空白 + 大小写归一 + 连字符归一，用于页眉页脚去重与公式占位符匹配。"""
    return _norm_line_from_strip(_strip_spaces(line))


_PAGE_NUMBER_EDGE_RE = re.compile(
    r"^\s*(?:p(?:age)?\.?\s*)?\d{1,4}(?:\s*[-–—/]\s*\d{1,4})?\s*$", re.IGNORECASE
)
_PAGE_NUMBER_DASHED_RE = re.compile(r"^\s*[-–—·]\s*\d{1,4}\s*[-–—·]\s*$")


def _is_page_number_edge(line: str) -> bool:
    if _PAGE_NUMBER_DASHED_RE.match(line):
        return True
    return bool(_PAGE_NUMBER_EDGE_RE.match(line)) and bool(re.search(r"\d", line))


# 中文期刊常见页眉/元数据行模式（用于辅助识别，即使不满足阈值也剔除）
_JOURNAL_HEADER_PATTERNS = [
    re.compile(r"综\s*述", re.IGNORECASE),
    re.compile(r"第\s*\d+\s*卷", re.IGNORECASE),
    re.compile(r"中国激光", re.IGNORECASE),
    re.compile(r"收稿日期", re.IGNORECASE),
    re.compile(r"修回日期", re.IGNORECASE),
    re.compile(r"录用日期", re.IGNORECASE),
    re.compile(r"网络首发日期", re.IGNORECASE),
    re.compile(r"基金项目", re.IGNORECASE),
    re.compile(r"通信作者", re.IGNORECASE),
    re.compile(r"DOI\s*[：:]\s*10\.", re.IGNORECASE),
    re.compile(r"中图分类号", re.IGNORECASE),
    re.compile(r"文献标志码", re.IGNORECASE),
    re.compile(r"^\d{7,}-\d{1,3}$", re.IGNORECASE),  # 如 "0907105-1"
    re.compile(r"^\d{7,}\s*$", re.IGNORECASE),  # 如 "0907105"（独立页码行）
]


def _looks_like_journal_metadata(line: str) -> bool:
    """判断一行是否像中文期刊的页眉/元数据行。"""
    stripped = line.strip()
    if not stripped:
        return False
    for pat in _JOURNAL_HEADER_PATTERNS:
        if pat.search(stripped):
            return True
    return False


def _strip_repeated_header_footer(pages: list[str]) -> list[str]:
    """剔除跨页重复出现的页眉/页脚行，并去掉页首/页尾的独立页码。

    页眉页脚（期刊名、卷期、DOI 链接、作者行等）通常出现在 >=30% 的页面，
    按「归一化后跨页出现次数」识别重复行。同时用中文期刊页眉模式做辅助匹配。
    页码行（如 ``12``、``- 5 -``）单独按形态剔除。
    """
    n = len(pages)
    if n < 1:
        return pages

    counts: dict[str, int] = {}
    for page in pages:
        seen: set[str] = set()
        for raw_line in page.splitlines():
            key = _norm_line(raw_line)
            if not key or key in seen:
                continue
            seen.add(key)
            counts[key] = counts.get(key, 0) + 1

    threshold = max(2, math.ceil(n * 0.3))
    repeated = {key for key, count in counts.items() if count >= threshold}

    cleaned: list[str] = []
    for page in pages:
        lines = [
            ln
            for ln in page.splitlines()
            if _norm_line(ln) not in repeated and not _looks_like_journal_metadata(ln)
        ]
        while lines and _is_page_number_edge(lines[0]):
            lines.pop(0)
        while lines and _is_page_number_edge(lines[-1]):
            lines.pop()
        cleaned.append("\n".join(lines).strip())
    return cleaned


# ---- 章节标题抽取（PDF 大纲优先，字号启发式兜底）----

_HEADING_CUE_RE = re.compile(
    r"(?i)^(?:\d{1,3}(?:[.\-–]\d{1,3}){0,2}[.\-–、\s]"
    r"|第[一二三四五六七八九十百\d]+[章节部分篇]"
    r"|chapter\s*\d+|section\s*\d+"
    r")\s*"
)
_BARE_HEADING_RE = re.compile(
    r"(?i)^(?:abstract|introduction|related work|methods?|experiments?|results?|discussion|conclusion|references?|acknowledgements?|"
    r"摘要|引言|绪论|前言|方法|实验|结果|讨论|结论|参考文献|致谢|附录)\s*$"
)
# 中文论文章节标题模式：数字编号 + 全角/半角空格 + 中文标题（必须含至少一个中文字符）
_CN_HEADING_RE = re.compile(
    r"^\s*(\d{1,2}(?:\.\d{1,2}){0,2})\s*[　\s]{1,6}\s*([\u4e00-\u9fff]+.*)"  # 如 "1　引　　言" "2.1　OCT 原理"
)
_SENTENCE_END_RE = re.compile(r"[。．.！？!?;；:：]\s*$")

# ---- 「像正文而不是像标题」的反例判据（只用于低置信标题候选）----
# 以虚词/连接词收尾 = 句子还没说完（「Blackbody：能吸收辐射到其」）。
_HEADING_STOP_TAIL = "的了和与及在是为其而但并或则把被从对将所能会可若如因由使得地着过"
# 标题编号前缀（「3.1 引言」「A. 材料」「1) 方法」）。必须带结尾空白，
# 否则会把 Introduction 的 I 当成罗马数字编号剥掉，把正文改成另一根字符串。
_HEADING_NUMBER_PREFIX_RE = re.compile(
    r"^\s*(?:\d{1,3}(?:[.\-–]\d{1,3}){0,2}|[A-Za-z]|[IVXLCDM]{1,6})[.\-–、)）:：]?\s+"
)


# 结构编号（「3.1」「第3章」「Chapter 2」）：几乎不可能是正文行，豁免反例判据。
# 裸序号（「1 引言」「2 实验」）不豁免 —— 幻灯片正文的编号列表项也长这样。
_HEADING_STRONG_CUE_RE = re.compile(
    r"(?i)^(?:\d{1,3}(?:[.\-–]\d{1,3})+[.\-–、\s]"
    r"|第[一二三四五六七八九十百\d]+[章节部分篇]"
    r"|chapter\s*\d+|section\s*\d+)"
)

_PLOT_LABEL_RE = re.compile(
    r"(?i)^\s*\d+(?:\.\d+)?\s*(?:[�Ωμµ×±−-]{1,4}\s*)"
    r"(?:load|input|output|signal|power|current|voltage|frequency|wavelength)\b"
    r"|^\s*\d+(?:\.\d+)?\s*(?:ghz|mhz|khz|thz|nm|μm|µm|db|mv|kw|mw|uw|ps|fs|v|w|a)\b"
)

# 图内坐标轴/图例常被 PDF 文本层拆成「数字 + 单位/短标签」的独立行，
# 也可能夹着乱码替换符（例如 ``50 �� load``）。这类行即使被字号兜底
# 识别成标题，也不应出现在目录里。只匹配整行，避免误删真正的长节标题
# （例如「3.95 MHz response」）。
_STANDALONE_PLOT_LABEL_RE = re.compile(
    r"(?ix)^\s*[+\-−]?\d+(?:\.\d+)?\s*"
    r"(?:[�Ωμµ×±−-]{0,4}\s*"
    r"(?:load|input|output|signal|power|current|voltage|frequency|wavelength)"
    r"|(?:ghz|mhz|khz|thz|nm|μm|µm|db|mv|kw|mw|uw|ps|fs|v|w|a))\s*$"
)
_PAGE_COUNT_LABEL_RE = re.compile(r"^\s*\d+\s*(?:of|/)\s*\d+\s*$", re.IGNORECASE)
_NUMERIC_AXIS_LABEL_RE = re.compile(
    r"(?ix)^\s*[+\-−]?\d+(?:\.\d+)?"
    r"(?:\s*(?:~|～|−|–|—|-|to)\s*[+\-−]?\d+(?:\.\d+)?)?"
    r"\s*(?:[×x*]\s*[+\-−]?\d+(?:\.\d+)?\s*)?"
    r"(?:[eE][+\-]?\d+\s*)?"
    r"(?:ghz|mhz|khz|thz|hz|nm|μm|µm|km|cm|mm|um|db|mv|kv|kw|mw|uw|nw|ps|ns|us|μs|µs|ms|s|m|v|w|a|k|%|deg)\s*$"
)
_DIMENSION_LABEL_RE = re.compile(
    r"^\s*\d+(?:\.\d+)?\s*[×x*]\s*\d+(?:\.\d+)?(?:\s*(?:px|pt|mm|cm|μm|µm|um))?\s*$",
    re.IGNORECASE,
)


def _looks_like_plot_label(text: str) -> bool:
    """坐标/图内标注（如 ``50 Ω load``、``10 GHz``）不是正文目录标题。"""
    value = str(text or "").strip()
    return bool(_PLOT_LABEL_RE.match(value) or _NUMERIC_AXIS_LABEL_RE.fullmatch(value))


def _looks_like_heading_noise(text: str) -> bool:
    """标题候选最终收口时过滤图内标签和页码计数。"""
    value = str(text or "").strip()
    return bool(
        _STANDALONE_PLOT_LABEL_RE.fullmatch(value)
        or _PAGE_COUNT_LABEL_RE.fullmatch(value)
        or _NUMERIC_AXIS_LABEL_RE.fullmatch(value)
        or _DIMENSION_LABEL_RE.fullmatch(value)
    )


def _heading_cue_strength(text: str) -> int:
    """标题编号特征的强度：2=结构编号/章节词，1=裸序号，0=无特征。"""
    if _HEADING_STRONG_CUE_RE.match(text) or _BARE_HEADING_RE.match(text):
        return 2
    return 1 if _HEADING_CUE_RE.match(text) else 0


def _looks_like_prose_line(text: str) -> bool:
    """判断一行是「正文句子」而不是标题。

    只作为**低置信候选的否决判据**使用：有编号特征、或字号明显大于正文的候选
    不走这里，避免把真标题拦掉。四类反例都来自实测的误判样本。
    """
    stripped = text.strip()
    if not stripped:
        return True
    # 剥掉编号前缀后再判：否则「3. 方法」会被下面的句号规则误杀
    body = _HEADING_NUMBER_PREFIX_RE.sub("", stripped).strip()
    if not body:
        return True
    # 句中标点：标题不在中间断句
    if re.search(r"[。！？!?；;]", body):
        return True
    # 句号后面不是数字/目录点号（「是相当的. 这」「运动.具有确定动量…」）；
    # 「2.1」「§1.1」的句号后面紧跟数字，不受影响
    if re.search(r"\.(?!\d)", body):
        return True
    # 逗号成句：正文句最典型的特征（「19世纪末，物理学理论…」）。
    # 中文标题用顿号/空格枚举，几乎不出现全角逗号 → 全角逗号门槛放低；
    # 英文标题形如「Methods, Materials and Results」确实带半角逗号，门槛放高，
    # 且这类标题通常还带编号（带编号即豁免本判据）。
    if len(body) >= 10 and re.search(r"，", body):
        return True
    if len(body) >= 20 and re.search(r",", body):
        return True
    # 冒号后接完整解释通常是幻灯片编号清单/正文句（如「1. 目的：压缩脉冲宽度」），
    # 而不是章节标题。这里只拦裸编号清单；结构编号（3.1、Chapter 2 等）以及
    # 无编号的真实标题（如「Methods: Experimental setup」）保留。
    if (
        len(body) >= 12
        and re.search(r"[:：]", body)
        and re.match(r"^\s*(?:\d{1,3}[.)、]|[（(]\d+[）)])\s+", stripped)
    ):
        return True
    # 截断式收尾
    if body[-1] in _HEADING_STOP_TAIL:
        return True
    return False


def _outline_headings(doc: Any) -> dict[int, list[tuple[int, str]]]:
    """从 PDF 书签提取章节标题，返回 {页码: [(level, 标题), ...]}。"""
    headings: dict[int, list[tuple[int, str]]] = {}
    try:
        toc = doc.get_toc(simple=True)
    except Exception:
        return headings
    for entry in toc or []:
        if not isinstance(entry, (list, tuple)) or len(entry) < 3:
            continue
        try:
            level = max(1, int(entry[0]))
        except (TypeError, ValueError):
            level = 1
        title = str(entry[1] or "").strip()
        if not title:
            continue
        try:
            page = int(entry[2])
        except (TypeError, ValueError):
            continue
        headings.setdefault(page, []).append((level, title))
    return headings


def _median(values: list[float]) -> float:
    if not values:
        return 0.0
    ordered = sorted(values)
    n = len(ordered)
    mid = n // 2
    return ordered[mid] if n % 2 else (ordered[mid - 1] + ordered[mid]) / 2.0


def _cn_heading_level(number_str: str) -> int:
    """根据中文论文章节编号推断层级：1→2级, 1.1→3级, 1.1.1→4级。"""
    dots = number_str.count(".")
    return min(4, 2 + dots)


# 公式行特征：大量数学符号（希腊字母、运算符、上下标、括号等）
_FORMULA_LINE_MATH_CHARS = re.compile(
    r"""[=≈≤≥<>≠≡±×÷⋅∑∫√∞∂∆∇ħℏα-ωΑ-Ωμθλπρσφψω\[\]{}|+*/^_]|\b(?:frac|sqrt|text|alpha|beta|gamma|delta|epsilon|zeta|eta|theta|iota|kappa|lambda|mu|nu|xi|pi|rho|sigma|tau|upsilon|phi|chi|psi|omega|sum|int|prod|coprod|partial|nabla|exp|ln|log|sin|cos|tan|det|lim)\b|-exp\b"""
)


def _looks_like_formula_line(text: str) -> bool:
    """判断一行是否像独立的公式行（非标题），避免误标记为 [[H:n]]。"""
    stripped = text.strip()
    if len(stripped) < 2:
        return False
    # 显式占位符或 LaTeX 块必为公式
    if "[[MATH:" in stripped or stripped.startswith("$$") or stripped.startswith("\\[") or "\\begin{" in stripped:
        return True
    has_chinese = bool(re.search(r"[\u4e00-\u9fff]", stripped))
    math_count = len(_FORMULA_LINE_MATH_CHARS.findall(stripped))
    has_parens = bool(re.search(r"[()（）]", stripped))
    has_equals = bool(re.search(r"[=≈≤≥≠≡∝∼]", stripped))
    has_formula_number = bool(re.search(r"[\(（]\s*\d{1,3}[a-zA-Z]?\s*[\)）]\s*$", stripped))

    # 含中文但公式特征明显（如 "F (w)，其表达式为" 或带等号/编号短行）
    if has_chinese:
        if has_equals and len(stripped) <= 60:
            return True
        if has_formula_number and len(stripped) <= 60:
            return True
        if len(stripped) <= 50 and has_parens and math_count >= 2:
            return True
        if len(stripped) <= 50 and re.search(r"[∑∫∏⋅βΔαμ∇∂]", stripped):
            return True
        return False

    # 无中文且含公式尾部编号（如 "... (1)"）→ 绝对是公式行
    if has_formula_number and (has_equals or math_count >= 1):
        return True
    # 无中文且含数学等号/不等号 → 纯算式，绝非标题
    if has_equals:
        return True
    # 无中文且含 2 个以上数学运算/函数/希腊字母符号
    if math_count >= 2:
        return True
    # 含微积分/积分/求和/绝对值等强数学符号
    if re.search(r"[∫∑∏∮√∞∂∇]|\|\||\|", stripped):
        return True
    return False


def _font_heading_lines(
    page: Any, page_number: int | None = None
) -> list[tuple[int, str, int | None]]:
    """按字号 + 加粗 + 短行启发式识别章标题，用于无书签 PDF 的兜底。

    `page_number`（从 1 起）随每条候选带回，最终写进 [[H:level@page]] 标记 ——
    目录的页码必须来自解析器，模型看不到页码，猜出来的必然是编造值。
    """
    try:
        raw = page.get_text("dict")
    except Exception:
        return []
    lines: list[dict[str, Any]] = []
    for block in raw.get("blocks", []):
        if block.get("type") != 0:
            continue
        for line in block.get("lines", []):
            spans = line.get("spans", [])
            if not spans:
                continue
            text = "".join(str(s.get("text", "")) for s in spans).strip()
            if not text:
                continue
            size = max((float(s.get("size", 0) or 0) for s in spans), default=0.0)
            bold = any(bool((s.get("flags", 0) or 0) & 16) for s in spans)
            italic = any(bool((s.get("flags", 0) or 0) & 2) for s in spans)
            lines.append(
                {
                    "text": text,
                    "size": size,
                    "bold": bold,
                    "italic": italic,
                    "n": len(text),
                    "bbox": [float(v) for v in line.get("bbox", (0, 0, 0, 0))],
                }
            )
    if not lines:
        return []

    long = [ln["size"] for ln in lines if ln["n"] >= 16]
    body = _median(long) if long else _median([ln["size"] for ln in lines])
    if body <= 0:
        return []

    headings: list[tuple[int, str]] = []
    for ln in lines:
        text = ln["text"]
        if ln["n"] > 80 or _SENTENCE_END_RE.search(text):
            continue
        if _looks_like_plot_label(text):
            continue
        # 跳过公式行：含大量数学符号但无中文字符
        if _looks_like_formula_line(text):
            continue
        # 跳过公式碎片行（如 "1-exp( 1-exp"），防止被误标为标题
        if _is_formula_debris_span(text):
            continue
        # 跳过单字符碎片
        if ln["n"] <= 1:
            continue
        # 跳过斜体行（公式通常使用斜体，非标题特征）
        if ln["italic"] and not re.search(r"[\u4e00-\u9fff]", text):
            continue
        # 跳过含空括号占位符的行（公式缺字形残留，如 "I [ ]i 和"）
        if re.search(r"[\[［]\s*[\]］]", text):
            continue
        # 跳过纯算式碎片（如 "1-exp("、"1 - exp["）：无中文且字母词全是数学函数名
        if not re.search(r"[\u4e00-\u9fff]", text):
            words = [w.lower() for w in re.findall(r"[A-Za-z]{2,}", text)]
            if (
                words
                and all(w in _DEBRIS_WORD_WHITELIST for w in words)
                and re.search(r"[\d=+\-*/()\[\]^]", text)
            ):
                continue
        ratio = ln["size"] / body
        # 中文论文章节标题：数字编号 + 空格 + 中文标题（排除公式特征）
        cn_match = _CN_HEADING_RE.match(text)
        if cn_match and ln["n"] <= 50 and not _looks_like_formula_line(text):
            # 编号清单项也长成「1 散射光中，除了…」的样子，同样要过正文反例判据
            if not _looks_like_prose_line(text):
                level = _cn_heading_level(cn_match.group(1))
                headings.append((level, text, ln))
                continue
        if not (ratio >= 1.18 or (ln["bold"] and ln["n"] <= 40)):
            continue
        # 无结构编号的候选（含只有裸序号的）都过一遍正文反例判据。
        # 带编号不等于不是正文，幻灯片上的编号列表项（实测：「1 散射光中，除了
        # 原来X光的波长λ外，增加了」，20pt、比正文中位数还大一倍）同样以
        # 「编号+空格」开头 —— 所以这里不能拿字号大小当免检牌。
        # 「3.1」「第2章」「Chapter 2」这类结构编号才免检。
        if _heading_cue_strength(text) <= 1:
            if ln["n"] <= 3:
                # 极短行只有两种可能：①换行折断的标题尾巴（「…计算精度的影响」+
                # 「因素」）②坐标轴刻/正文碎片（「波长」「所以」「4.2」）。一律标成
                # 「碎片候选」：允许被 _merge_wrapped_headings 并进上一行标题，
                # 但绝不允许自己成为一条目录项。
                if re.fullmatch(r"[\u4e00-\u9fff]{1,3}", text):
                    ln["_fragment"] = True
                    headings.append((3, text, ln))
                continue
            if _looks_like_prose_line(text):
                continue
        # 严格过滤公式行、公式碎片与带数学符号/编号的非标题行
        if _looks_like_formula_line(text) or _is_formula_debris_span(text):
            continue
        if re.search(r"[=≈≤≥≠≡∝∼^_{}\\]|[\(（]\s*\d{1,3}\s*[\)）]\s*$", text):
            continue
        # 首页非编号行要求更大字号：过滤作者/单位等元数据行（通常仅略大于正文）
        if (
            getattr(page, "number", 1) == 0
            and not _HEADING_CUE_RE.match(text)
            and not _BARE_HEADING_RE.match(text)
            and ratio < 1.35
        ):
            continue
        # 无中文行必须具备明确的标题特征（匹配预设词或编号英文标题），绝不放行纯大字号英文/公式碎片
        has_cn = bool(re.search(r"[\u4e00-\u9fff]", text))
        if not has_cn:
            if not (_HEADING_CUE_RE.match(text) or _BARE_HEADING_RE.match(text) or re.match(r"^(?:\d+(?:\.\d+)*|[IVXLCDM]+\.)\s+[A-Za-z]", text)):
                continue
        elif not (_HEADING_CUE_RE.match(text) or _BARE_HEADING_RE.match(text) or ln["n"] <= 40):
            continue
        # level 语义与前端/ MinerU 对齐：2=节、3=小节（1 保留给论文题名，此处不产出）
        level = 2 if ratio >= 1.35 else 3
        headings.append((level, ln["text"], ln))
    return _merge_wrapped_headings(headings, page_number)


def _merge_wrapped_headings(
    cands: list[tuple[int, str, dict[str, Any]]],
    page_number: int | None = None,
) -> list[tuple[int, str, int | None]]:
    """合并换行折断的编号标题（如 "4　组织光学衰减系数计算精度的影响" + "因素"）。

    下一条候选须是同字号、紧邻下方、左对齐、无编号无标题特征的短中文碎片。
    全部候选来自同一页，页码统一带出。
    """
    used = [False] * len(cands)
    merged: list[tuple[int, str, int | None]] = []
    for i, (level, text, ln) in enumerate(cands):
        # 碎片候选只作为合并对象，自己不成条（否则「波长」「所以」会进目录）
        if used[i] or ln.get("_fragment"):
            continue
        if _CN_HEADING_RE.match(text):
            for j in range(i + 1, len(cands)):
                if used[j]:
                    continue
                _, t2, ln2 = cands[j]
                t2 = t2.strip()
                if (
                    len(t2) <= 8
                    and re.fullmatch(r"[\u4e00-\u9fff、·]+", t2)
                    and not _CN_HEADING_RE.match(t2)
                    and not _HEADING_CUE_RE.match(t2)
                    and not _BARE_HEADING_RE.match(t2)
                    and abs(ln2["size"] - ln["size"]) <= 0.6
                    and -2.0 <= ln2["bbox"][1] - ln["bbox"][3] <= 12.0
                    and abs(ln2["bbox"][0] - ln["bbox"][0]) <= 40.0
                    and len(text) + len(t2) <= 50
                ):
                    text = text.rstrip() + t2
                    used[j] = True
                    break
        merged.append((level, text.strip(), page_number))
    return merged


def _normalize_heading_levels(
    headings: list[tuple[Any, ...]],
) -> list[tuple[Any, ...]]:
    """把 PDF 大纲书签的层级平移，使最浅一级映射为 2（节）；1 留给题名。

    入参可以是 (level, title) 或 (level, title, page)，页码原样带过。
    """
    if not headings:
        return headings
    min_level = min(int(item[0]) for item in headings)
    shift = 2 - min_level
    if shift == 0:
        return headings
    return [(max(2, min(4, int(item[0]) + shift)), *item[1:]) for item in headings]


def _dedupe_headings(headings: list[tuple[Any, ...]]) -> list[tuple[int, str, int | None]]:
    """按归一化标题去重（大纲与字号启发式可能重复识别同一条标题）。

    统一收敛成 (level, title, page) 三元组，page 未知时为 None。
    """
    seen: set[str] = set()
    out: list[tuple[int, str, int | None]] = []
    for item in headings:
        level, title = item[0], item[1]
        page = item[2] if len(item) > 2 else None
        key = _norm_line(str(title))
        if not key or key in seen:
            continue
        seen.add(key)
        out.append((int(level), str(title), page))
    return out


def _heading_marker(level: int, page: Any = None) -> str:
    """标题标记：带真实页码时为 [[H:level@page]]，无页码退回 [[H:level]]。"""
    try:
        page_number = int(page)
    except (TypeError, ValueError):
        page_number = 0
    if page_number > 0:
        return f"[[H:{int(level)}@{page_number}]]"
    return f"[[H:{int(level)}]]"


def _mark_headings_in_text(text: str, headings: list[tuple[Any, ...]]) -> str:
    """把独立成行的标题行替换为 [[H:level@page]] 前缀。"""
    lines = [ln for ln in text.split("\n")]
    used: set[int] = set()
    for item in headings:
        level, title = item[0], item[1]
        page = item[2] if len(item) > 2 else None
        marker = _heading_marker(level, page)
        key = _norm_line(str(title))
        if not key:
            continue
        # 安全防线：若标题候选本身像公式行，绝不作为标题打标进正文
        if (
            _looks_like_formula_line(title)
            or _is_formula_debris_span(title)
            or _looks_like_heading_noise(title)
        ):
            continue
        matched = False
        for idx, line in enumerate(lines):
            if idx in used:
                continue
            # 若行已包含公式/图片标记，或该行像公式行，绝不打上大标题标签
            if "[[MATH:" in line or "[[FIG:" in line or _looks_like_formula_line(line):
                continue
            if line.strip() and _norm_line(line) == key:
                lines[idx] = f"\n\n{marker} {line.strip()}\n\n"
                used.add(idx)
                matched = True
                break
        if not matched:
            # 标题在 PDF 里折成两行（如长节标题），尝试相邻两行拼接匹配
            for idx in range(len(lines) - 1):
                if idx in used or idx + 1 in used:
                    continue
                a, b = lines[idx].strip(), lines[idx + 1].strip()
                if "[[MATH:" in a or "[[MATH:" in b or _looks_like_formula_line(a) or _looks_like_formula_line(b):
                    continue
                if a and b and _norm_line(a) + _norm_line(b) == key:
                    sep = "" if re.search(r"[\u4e00-\u9fff]$", a) else " "
                    lines[idx] = f"\n\n{marker} {a}{sep}{b}\n\n"
                    lines[idx + 1] = ""
                    used.add(idx)
                    used.add(idx + 1)
                    matched = True
                    break
    return "\n".join(lines)


# 带标题但**没有页码**的标记：后处理补出来的（见 _reflow_paragraphs 的编号短标题兜底）
_HEADING_MARKER_NO_PAGE_RE = re.compile(r"(?m)^\[\[H:(\d)\]\] ([^\n]+)$")


def _attach_heading_pages(text: str, page_texts: list[str]) -> str:
    """给正文里补出来的 [[H:n]] 标记回填真实页码（改成 [[H:n@p]]）。

    页码只有解析阶段知道；`_reflow_paragraphs` 的「编号短标题兜底」是在整篇文本上
    补标记的，天然没有页码，前端目录因此少了页码。按「这条标题行落在哪一页」回填。
    """
    if not page_texts or "[[H:" not in text:
        return text
    page_line_index: list[dict[str, int]] = []
    for index, page_text in enumerate(page_texts):
        lines: dict[str, int] = {}
        for line in page_text.split("\n"):
            key = _norm_line(line)
            if key:
                lines.setdefault(key, index + 1)
        page_line_index.append(lines)

    def _sub(match: re.Match[str]) -> str:
        level, title = match.group(1), match.group(2)
        key = _norm_line(title)
        if not key:
            return match.group(0)
        for lines in page_line_index:
            page = lines.get(key)
            if page:
                return f"[[H:{level}@{page}]] {title}"
        return match.group(0)

    return _HEADING_MARKER_NO_PAGE_RE.sub(_sub, text)


# 独立成行、**没有标题文本**的标题标记。正常标记长这样：`[[H:3@7]] 3.1 单散射模型`。
# 空标记来自后处理（公式候选还原、碎片清行）把标题文本挪到了别处。
_NAKED_HEADING_MARKER_RE = re.compile(r"(?m)^[ \t]*\[\[H:\d+(?:@\d+)?\]\][ \t]*$")


def _normalize_heading_markers(text: str) -> str:
    """清掉空标记以及明显的图内标签/页码伪标题。

    留着这类标记有两个可见后果：阅读器渲染出一个空标题块；前端目录的正则
    会把紧随其后的正文行当成标题文本（「所以」「波长」「4.2」这类目录项就是这么来的）。
    带标题的标记不受影响。
    """
    if "[[H:" not in text:
        return text
    cleaned = _NAKED_HEADING_MARKER_RE.sub("", text)

    # 目录候选可能来自 `_reflow_paragraphs` 的编号短标题兜底，无法携带
    # 字号/坐标信息；在这里做最后一道与来源无关的形态过滤，避免诸如
    # ``[[H:3]] 3.95 MHz``、``[[H:3]] 6 of 6`` 泄漏到前端目录。
    heading_line_re = re.compile(
        r"(?m)^[ \t]*\[\[H:\d+(?:@\d+)?\]\][ \t]+([^\n]+)[ \t]*$"
    )

    def _drop_noise_marker(match: re.Match[str]) -> str:
        title = match.group(1).strip()
        return "" if _looks_like_heading_noise(title) else match.group(0)

    cleaned = heading_line_re.sub(_drop_noise_marker, cleaned)
    if cleaned == text:
        return text
    return re.sub(r"\n{3,}", "\n\n", cleaned)


_HEADING_LIKE_RE = re.compile(
    r"^(\d+(?:\.\d+){0,3})[　\s]+([^\d。，；！？.,;!?][^。，；！？]{0,28})$"
)
# 独立成行的邮箱（页脚通信作者信箱，可带 * 上标注释）
_EMAIL_LINE_RE = re.compile(r"^\s*\*{0,2}[\w.+-]+@[\w-]+(?:\.[\w-]+)+\s*$")


def _is_heading_like(line: str) -> bool:
    """标题形行：[[H:]] 标记行，或「2　技术背景」「2.1　OCT 原理」式编号短标题。"""
    s = line.strip()
    if not s:
        return False
    if s.lstrip().startswith("[[H:"):
        return True
    return bool(_HEADING_LIKE_RE.fullmatch(s))


def _join_wrapped_lines(lines: list[str]) -> str:
    """把 PDF 按视觉行硬切出来的多个短行拼回一个自然段，顺带还原英文跨行连字符。"""
    parts: list[str] = []
    for line in lines:
        if not parts:
            parts.append(line)
            continue
        prev = parts[-1]
        # 标题行不参与段落拼接：否则「2　技术背景」「2.1　OCT 原理」会被
        # 中文连续规则吞进正文段（[[H:2]] 2　技术背景2.1　OCT 原理 OCT 的基本原理是…）
        if _is_heading_like(line) or _is_heading_like(prev):
            parts.append(line)
            continue
        # 英文断词：上一行以「字母/数字 + -」结尾，下一行以小写字母开头 → 拼回同一个词
        if re.fullmatch(r".*[A-Za-z0-9]-\Z", prev) and re.fullmatch(r"[a-z].*", line):
            parts[-1] = prev[:-1] + line
            continue
        # 中文连续：上一行以中文字符结尾且下一行以中文字符开头，同句被截断 → 直接拼接
        if re.search(r"[\u4e00-\u9fff\uff00-\uffef]\s*$", prev) and re.match(
            r"^\s*[\u4e00-\u9fff\uff00-\uffef]", line
        ):
            parts[-1] = prev + line
            continue
        parts.append(line)
    return " ".join(parts)


_FORMULA_MARKER_LINE_RE = re.compile(
    r"^\s*\[\[(?:MATH|FIG):[^\]]+\]\]\s*$"
)
_TERMINAL_PUNCT = "。；？！…："
_CJK_RE = re.compile(r"[\u4e00-\u9fff]")
_CJK_TAIL_RE = re.compile(r"[\u4e00-\u9fff\)\)\]）】」』]$")


def _reflow_paragraphs(text: str) -> str:
    """把单换行合并为空格、保留空行（段落）边界，让正文按自然段流动显示。

    跨段合并：相邻块在「上一块末尾是未终止的 CJK」+「下一块开头是 CJK 或 inline 公式字符」
    时合并——处理跨页断句（如「OCT 系\n\\n统的有限探测深度」被页面边界硬切）。
    合并迭代直到稳定，避免三个块被一次推平。
    """
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = text.replace("■", "").replace("\ufffd", "")
    blocks = re.split(r"\n[ \t]*\n", text)
    paragraphs: list[str] = []
    for block in blocks:
        lines = [ln.strip() for ln in block.split("\n")]
        lines = [ln for ln in lines if ln]
        if not lines:
            continue
        # 公式/图片标记行、标题行独立成段：否则段落重排会把标记/标题前后的
        # 半句粘成「…[[MATH:f30]]…」夹心句或「2　技术背景2.1　OCT 原理 OCT 的…」粘连段
        parts: list[list[str]] = [[]]
        for ln in lines:
            if _FORMULA_MARKER_LINE_RE.fullmatch(ln):
                parts.append([ln])
                parts.append([])
            elif _is_heading_like(ln):
                parts.append([ln])
                parts.append([])
            else:
                parts[-1].append(ln)
        for part in parts:
            part = [ln for ln in part if ln]
            if not part:
                continue
            if len(part) == 1 and _FORMULA_MARKER_LINE_RE.fullmatch(part[0]):
                paragraphs.append(part[0])
                continue
            # 无标记的编号短标题行（如「2.1　OCT 原理」漏了字体识别）→ 独立成段并补 [[H:n]]
            if len(part) == 1:
                m = _HEADING_LIKE_RE.fullmatch(part[0].strip())
                if m and not part[0].lstrip().startswith("[[H:"):
                    tail = m.group(2).strip()
                    # 「编号 + 空白 + 短尾巴」不等于标题：列表项/公式碎片也长这样
                    # （实测误判样本：「0 )」「2 波长增量 Δλ=λ’–λ 随散射角增大而增大」）。
                    # 要求尾巴是像样的标题文本。
                    if (
                        len(tail) >= 2
                        # 中文尾（「引言」「OCT 原理」）或首字母大写的英文尾
                        # （「Materials and Methods」）。纯小写英文/单位尾（「mm」
                        # 「eyes from」）是坐标轴刻度与正文碎片，不是标题。
                        and (
                            re.search(r"[\u4e00-\u9fff]", tail)
                            or (len(tail) >= 3 and tail[:1].isupper())
                        )
                        and not _looks_like_formula_line(part[0])
                        and not _looks_like_heading_noise(part[0])
                        and not _looks_like_prose_line(tail)
                    ):
                        level = m.group(1).count(".") + 2
                        paragraphs.append(f"[[H:{level}]] {part[0].strip()}")
                        continue
            paragraphs.append(_join_wrapped_lines(part))

    # 跨段合并：上一段「未终止的 CJK 收尾」+ 下一段「CJK/inline 公式字符开头」→ 拼接到上段
    def _continuable(prev: str, cur: str) -> bool:
        if not prev or not cur:
            return False
        if _FORMULA_MARKER_LINE_RE.fullmatch(cur) or cur.startswith("[[H:"):
            return False
        # 标题段自身也不吞并后续正文（否则「1　引　　言」+ 首段正文被拼回一段）
        if prev.startswith("[[H:"):
            return False
        last = prev[-1]
        if last in _TERMINAL_PUNCT or last in ",.;:!?\"'":
            return False
        # 上一段末是 CJK 才认为可能跨段续写（避免合并英文段落或列表）
        if not _CJK_TAIL_RE.search(prev):
            return False
        # 下一段首是 CJK、inline 公式字符（数字/希腊/算符）才接续
        first = cur.lstrip()[0] if cur.lstrip() else ""
        if _CJK_RE.match(first) or first.isdigit() or first in "∑∫∏≈≅≡≪≫±×÷·⌈⌊√∛∜α-ωΑ-Ω":
            return True
        return False

    merged: list[str] = []
    for p in paragraphs:
        if merged and _continuable(merged[-1], p):
            merged[-1] = merged[-1] + p
        else:
            merged.append(p)
    return "\n\n".join(merged)


def _assign_formula_ids(formulas: list[dict[str, Any]]) -> list[dict[str, Any]]:
    for index, formula in enumerate(formulas):
        formula["id"] = f"f{index}"
    return formulas


def _replace_standalone_line(text: str, raw: str, marker: str) -> str:
    """把「整行即公式」的行替换为占位符（按整行匹配，不会误伤行内片段）。"""
    target = _norm_line(raw)
    if not target:
        return text
    out: list[str] = []
    for line in text.split("\n"):
        if line.strip() and _norm_line(line) == target:
            out.append(f"\n\n{marker}\n\n")
        else:
            out.append(line)
    return "\n".join(out)


def _mark_formulas_in_text(text: str, formulas: list[dict[str, Any]]) -> str:
    marked = text
    # 行归一结果按「行字符串」缓存。同一个公式要扫全部行，而每个公式都要扫一遍，
    # 不缓存就会做 O(公式数 × 行数) 次正则归一：300 条公式 × 2 万行的长文献实测
    # 单这一步就要 2~5 分钟，是「导入超时」的真实根因（解析本身只要十几秒）。
    line_strip_memo: dict[str, str] = {}
    for formula in sorted(
        formulas, key=lambda f: len(str(f.get("raw") or "")), reverse=True
    ):
        raw = str(formula.get("raw") or "").strip()
        if not raw:
            continue
        # raw 里混入占位标记（多行公式候选从已插标记的正文中捕获）：
        # 若照常替换会把别人的 [[MATH:]]/[[FIG:]] 吞进该公式的标记替换里，
        # 后续 _purge_bogus_formulas 还原 raw 时又把旧标记吐回正文（僵尸标记）
        if "[[" in raw:
            continue
        marker = f"[[MATH:{formula['id']}]]"
        raw_norm = _norm_line(raw)
        # 「含编号尾段」候选只取决于 raw、与具体行无关，整轮行扫描前算一次即可。
        # 原实现放在逐行调用的 _match_inline_in_line 内部重复生成，是最大的热点。
        raw_strip = _strip_spaces(raw)
        inline_candidates = _inline_tail_candidates(raw_strip)
        inline_enabled = len(raw_norm) >= 8 and _DIGIT_RE.search(raw) is not None
        lines = marked.split("\n")
        replaced = False
        for idx, line in enumerate(lines):
            line_strip = line_strip_memo.get(line)
            if line_strip is None:
                line_strip = _strip_spaces(line)
                line_strip_memo[line] = line_strip
            line_norm = _norm_line_from_strip(line_strip)
            # 整行精确匹配（display 公式场景）
            if line_norm == raw_norm:
                lines[idx] = f"\n\n{marker}\n\n"
                replaced = True
                break
            # 行尾带 (N) 时剥掉编号后比较
            if _FORMULA_TRAILING_NUM_RE.sub("", line_norm) == raw_norm:
                num_match = _FORMULA_TRAILING_NUM_RE.search(line_norm)
                if num_match and not formula.get("number"):
                    formula["number"] = num_match.group(1)
                lines[idx] = f"\n\n{marker}\n\n"
                replaced = True
                break
            # 行内公式：raw_norm 是 line_norm 的子串（pdfplumber 常把整行当 raw 捕获）
            # 或 raw 的尾段（含公式编号）能在 line 中找到——后者更常见，按"含编号尾段"匹配
            if inline_enabled:
                hit = _match_inline_in_line(
                    line, raw_strip, inline_candidates, line_strip
                )
                if hit is not None:
                    start, end = hit
                    lines[idx] = line[:start] + f" {marker} " + line[end:]
                    replaced = True
                    break
        if replaced:
            marked = "\n".join(lines)
        elif raw in marked:
            marked = marked.replace(raw, f"\n\n{marker}\n\n")

    return re.sub(r"\n{3,}", "\n\n", marked)


def _inline_tail_candidates(raw_strip: str) -> list[str]:
    """raw 的「含编号尾段」候选（12~24 字符），长尾段优先。

    只依赖 raw、与具体行无关，所以可以在整轮行扫描前算一次。
    """
    candidates: list[str] = []
    n = len(raw_strip)
    upper = min(n, 24)
    lower = min(12, n)
    for L in range(upper, lower - 1, -1):
        for start in range(0, n - L + 1):
            seg = raw_strip[start:start + L]
            if "（" in seg or "(" in seg or _DIGIT_RE.search(seg) is not None:
                candidates.append(seg)
    # 先试长尾段（含括号编号）
    return sorted(candidates, key=len, reverse=True)


def _match_inline_in_line(
    line: str,
    raw_strip: str,
    candidates: list[str] | None = None,
    line_strip: str | None = None,
) -> tuple[int, int] | None:
    """行内公式子串定位：在 line 中找 raw 的「含编号的最长尾段」。

    pdfplumber 等识别常把整行文本当公式 raw 捕获，raw 比 line 多含「前一行或后一行」
    残留。匹配策略：取 raw 去空格后的若干尾段（12-24 字符），含括号编号的最长者
    在 line 中出现即视为 inline 公式，标记为对应位置。

    candidates / line_strip 可由调用方预计算传入，避免逐行重复生成候选。
    """
    if not raw_strip:
        return None
    if line_strip is None:
        line_strip = _strip_spaces(line)
    if candidates is None:
        candidates = _inline_tail_candidates(raw_strip)
    # 优先：raw 整体在 line_strip 中
    if raw_strip in line_strip:
        return _find_in_line(line, raw_strip)
    # 备选：取 raw 的含编号尾段（12~24 字符），在 line_strip 中查找
    for tail in candidates:
        if tail in line_strip:
            return _find_in_line(line, tail)
    return None


def _find_in_line(line: str, raw: str) -> tuple[int, int] | None:
    """在 line 中定位 raw（去空格后）的字符区间。raw 比 line 短也能找到。"""
    raw_strip = re.sub(r"[\s\u00a0]+", "", raw)
    line_strip = re.sub(r"[\s\u00a0]+", "", line)
    p = line_strip.find(raw_strip)
    if p < 0:
        return None
    # 跳过 raw_strip 中可能与 line_strip 同名但实际跨字符的歧义，这里 raw 短而 line 长，
    # 不会有交叉；用 norm 长度差兜底：取 line 中去除 p 字符、保留 p+len(raw_strip) 字符
    # 之后按字符数推回原 line 索引。简单实现：按字符逐一扫描 line 跳过空白。
    i = 0
    pi = 0
    start = end = -1
    while i < len(line) and pi < p:
        if line[i].isspace():
            i += 1
            continue
        pi += 1
        i += 1
    start = i
    pi = 0
    while i < len(line) and pi < len(raw_strip):
        if line[i].isspace():
            i += 1
            continue
        pi += 1
        i += 1
    end = i
    return (start, end)


# —— 公式编号挂载与公式碎片清理 ——————————————————————————————

_MATH_MARKER_ANYWHERE_RE = re.compile(r"\[\[MATH:([^\]]+)\]\]")
_EQ_NUM_TOKEN_RE = re.compile(r"[\(（]\s*(\d{1,3}[a-zA-Z]?)\s*[\)）]")
_DEBRIS_SIGNAL_RE = re.compile(
    r"[=≈∝∫∑∏≥≤±\u2212\u2013]|⋅|·|\bexp\b|\bln\b|\bsin\b|\bcos\b|[μλγβσζρκθπαω]"
)
_DEBRIS_WORD_WHITELIST = {
    "exp", "ln", "log", "sin", "cos", "tan", "det", "lim", "const", "where", "if",
}


def _is_formula_debris_span(span: str, min_signals: int = 2) -> bool:
    """严格判定一段文本是否为公式碎片：基本无中文、数学信号密集、无正常英文单词。"""
    s = span.strip()
    if len(s) < 2:
        return False
    if len(re.findall(r"[一-鿿]", s)) > 2:
        return False
    if len(_DEBRIS_SIGNAL_RE.findall(s)) < min_signals:
        return False
    words = [w.lower() for w in re.findall(r"[A-Za-z]{3,}", s)]
    bad = [w for w in words if w not in _DEBRIS_WORD_WHITELIST]
    if len(bad) > 3:
        return False
    return True


def _attach_formula_numbers(text: str, formulas: list[dict[str, Any]]) -> str:
    """把游离在正文里的公式编号（N）挂到最近的 [[MATH:id]] 条目上，并从正文剥离。

    PDF 文本层里编号常与公式分离（独立行、下一行行首、碎片行尾），
    挂载后前端才能把编号渲染到公式右侧。「式（N）」是交叉引用，不动。
    """
    by_id = {str(f.get("id")): f for f in formulas or [] if f.get("id")}
    if not by_id or "[[MATH:" not in text:
        return text
    lines = text.split("\n")

    def attach(num: str, idx: int) -> bool:
        for f in by_id.values():
            if str(f.get("number") or "") == num:
                return False  # 编号已被占用
        for dist in range(1, 6):
            for j in (idx - dist, idx + dist):  # 先向前找（编号通常在公式后）
                if 0 <= j < len(lines):
                    m = _MATH_MARKER_ANYWHERE_RE.search(lines[j])
                    if m:
                        f = by_id.get(m.group(1))
                        if f is not None and not f.get("number"):
                            f["number"] = num
                            f["_num_src"] = "prox"  # 距离挂载：低置信，可被引用反推覆盖
                            return True
        return False

    out: list[str] = []
    for idx, line in enumerate(lines):
        stripped = line.strip()
        if not stripped:
            out.append(line)
            continue
        # 1) 整行只有编号
        m = re.fullmatch(r"[\(（]\s*(\d{1,3}[a-zA-Z]?)\s*[\)）]", stripped)
        if m:
            if attach(m.group(1), idx):
                continue
            out.append(line)
            continue
        # 2) 行首编号：「（7）然后…」——前提是本行不再含其它非引用的编号（排除枚举列表）。
        # 「（N）式中/式…」是交叉引用句，不剥离（编号由第二遍引用反推挂载）。
        m = re.match(r"^[\(（]\s*(\d{1,3}[a-zA-Z]?)\s*[\)）]\s*", line)
        if m:
            rest = line[m.end() :]
            other_nums = [
                t for t in _EQ_NUM_TOKEN_RE.finditer(rest)
                if rest[t.start() - 1 : t.start()] not in ("式", "见")
            ]
            if (
                not re.match(r"\s*式", rest)
                and not other_nums
                and attach(m.group(1), idx)
            ):
                line = rest
                stripped = line.strip()
                if not stripped:
                    continue
        # 3) 行尾编号（前一字符不是「式/见」）
        m = re.search(r"[\(（]\s*(\d{1,3}[a-zA-Z]?)\s*[\)）]\s*$", line)
        if m and line[m.start() - 1 : m.start()] not in ("式", "见") and attach(m.group(1), idx):
            line = line[: m.start()].rstrip()
            stripped = line.strip()
            if not stripped:
                continue
        else:
            # 4) 行中编号：后跟 对/而/则/标点（前一字符不是「式/见」）。
            # 注意「（N）式中」是交叉引用句，不能在此剥离（留给第二遍引用反推）。
            m = re.search(
                r"[\(（]\s*(\d{1,3}[a-zA-Z]?)\s*[\)）](?=\s*(?:对|而|则|，|,|。))",
                line,
            )
            if (
                m
                and line[m.start() - 1 : m.start()] not in ("式", "见")
                and attach(m.group(1), idx)
            ):
                line = (line[: m.start()] + line[m.end() :]).strip()
                if not line:
                    continue
        out.append(line)

    text = "\n".join(out)

    # 第二遍：用交叉引用反推仍缺失的编号（只挂载，正文引用原样保留）。
    # PDF 提取常把公式右侧编号丢失，但正文引用句能指出该公式是第几号。
    # - 「（N）式中：…」是解释句，公式必在其前 → 向前找最近无编号标记。
    # - 「…如式（N）所示」是预告句，公式通常紧跟其后 → 向后找最近无编号标记，
    #   且要求该标记后数行内出现同号「（N）式中」印证（排除「由式（N）可得」引出新公式的场景）。
    lines = text.split("\n")
    used = {str(f.get("number")) for f in by_id.values() if f.get("number")}
    ahead_pat = re.compile(r"[\(（]\s*(\d{1,3}[a-zA-Z]?)\s*[\)）](?=\s*式)")  # （N）式中：…
    trail_pat = re.compile(r"(?<=式)\s*[\(（]\s*(\d{1,3}[a-zA-Z]?)\s*[\)）]")  # 式（N）/ 如式（N）

    def math_marker_at(j: int) -> dict[str, Any] | None:
        mm = _MATH_MARKER_ANYWHERE_RE.search(lines[j])
        return by_id.get(mm.group(1)) if mm else None

    def attach_ref(num: str, j: int) -> bool:
        f = math_marker_at(j)
        if f is None:
            return False
        if not f.get("number"):
            f["number"] = num
            f["_num_src"] = "ref"
            used.add(num)
            return True
        # 引用句（高置信）与距离挂载（低置信）冲突时覆盖后者；
        # 被顶替的旧编号不复用（其正文出处已被剥离，无法再精确定位）。
        if f.get("_num_src") == "prox" and num not in used:
            f["number"] = num
            f["_num_src"] = "ref"
            used.add(num)
            return True
        return False

    # 先处理高置信的「（N）式中」解释句
    for idx, line in enumerate(lines):
        for m in ahead_pat.finditer(line):
            num = m.group(1)
            if num in used:
                continue
            for j in range(idx - 1, max(-1, idx - 13), -1):
                if _MATH_MARKER_ANYWHERE_RE.search(lines[j]):
                    attach_ref(num, j)
                    break
    # 再处理需印证的「式（N）」预告句
    for idx, line in enumerate(lines):
        for m in trail_pat.finditer(line):
            num = m.group(1)
            if num in used:
                continue
            for j in range(idx + 1, min(len(lines), idx + 13)):
                if _MATH_MARKER_ANYWHERE_RE.search(lines[j]):
                    confirm = any(
                        nm.group(1) == num
                        for k in range(j + 1, min(len(lines), j + 5))
                        for nm in ahead_pat.finditer(lines[k])
                    )
                    if confirm:
                        attach_ref(num, j)
                    break
    for f in by_id.values():
        f.pop("_num_src", None)
    return text


def _drop_formula_debris_lines(text: str) -> str:
    """清除 [[MATH:id]] 标记附近的公式碎片行（多行公式未被整体匹配时残留的半行）。

    三种形态：整行碎片直接删除；行首碎片前缀、行尾碎片后缀剥掉（保留同线的正文）。
    远离公式标记的行一律不动，避免误伤正文。
    """
    if "[[MATH:" not in text:
        return text
    lines = text.split("\n")
    marker_idx = [i for i, ln in enumerate(lines) if _MATH_MARKER_ANYWHERE_RE.search(ln)]
    if not marker_idx:
        return text

    # 每个公式标记所在的「段落块」（空行分隔）：块内残留的积分上下标/行内分数碎片
    # （"0"、"2"、"1"…）即使离标记 >±4 行也应清理；块外的正文段落不受影响
    block_id: list[int] = []
    bid = 0
    for ln in lines:
        if not ln.strip():
            bid += 1
        block_id.append(bid)
    marker_blocks = {block_id[m] for m in marker_idx}

    def near(i: int, window: int) -> bool:
        if any(abs(i - m) <= window for m in marker_idx):
            return True
        return block_id[i] in marker_blocks


    out: list[str] = []
    for i, line in enumerate(lines):
        stripped = line.strip()
        if not stripped or _MATH_MARKER_ANYWHERE_RE.search(stripped):
            out.append(line)
            continue
        # 与上一非空行重复的碎片行 → 去重（LLM 重排偶发逐字重复）
        prev_nonblank = next((x.strip() for x in reversed(out) if x.strip()), "")
        if (
            prev_nonblank
            and stripped == prev_nonblank
            and _is_formula_debris_span(stripped)
        ):
            continue
        if not near(i, 4):
            out.append(line)
            continue
        # 紧邻公式标记的 1-2 字符纯符号/变量行（")"、"|"、"r"、"1"…）：行内公式
        # 上下标残留，无中文即不可能是正文句子
        if re.fullmatch(r"[A-Za-z0-9+\-*/=()（）{}\[\]|^_.,:;⋅×÷<>≤≥≈∝∫∑∏√±∞'′°πΠμΜθΘλΛωΩψΨφΦσΣκΚβΒαΑγΓζΖρΡνΝξΞτΤδΔεΕ]{1,2}", stripped) and abs(i - min(marker_idx, key=lambda m: abs(m - i))) <= 2:
            continue
        # 上标/下标数字碎片（公式区域图片化后残留的 "2"、"2 2" 等；编号（N）带括号不受影响）
        if re.fullmatch(r"(?:\d\s*){1,6}", stripped):
            continue
        # 整行碎片
        if len(stripped) <= 120 and _is_formula_debris_span(stripped):
            continue
        # 行首碎片前缀（后接中文正文）
        m = re.match(r"^([^\u4e00-\u9fff。；，：、？！]{4,60}?)(?=[\u4e00-\u9fff])", stripped)
        if m and _is_formula_debris_span(m.group(1)):
            stripped = stripped[m.end() :]
            line = stripped
            if not stripped:
                continue
        # 行尾碎片后缀（最后一个中文/中文标点之后、≥3 个数学信号）
        m = re.search(r"[\u4e00-\u9fff。；，：、？！]([^\u4e00-\u9fff。；，：、？！]{6,})$", line)
        if m and _is_formula_debris_span(m.group(1), min_signals=3):
            line = line[: m.start(1)].rstrip(" \t，,：:")
            if not line.strip():
                continue
        out.append(line)
    return re.sub(r"\n{3,}", "\n\n", "\n".join(out))


def _purge_bogus_formulas(text: str, formulas: list[dict[str, Any]]) -> str:
    """剔除误判为公式的正文片段（raw 含 ≥3 个中文字符）：还原其占位符为原文。"""
    kept: list[dict[str, Any]] = []
    changed = False
    for f in formulas or []:
        raw = str(f.get("raw") or f.get("text") or "")
        if len(re.findall(r"[\u4e00-\u9fff]", raw)) >= 3:
            marker = f"[[MATH:{f.get('id')}]]"
            # 这条「公式」其实是标题行上的正文（如「[[H:3]] 中只能包含ψ, …」被判成
            # 公式候选）。_mark_formulas_in_text 会把标题文本整段换成公式标记，
            # 还原时必须挂回 [[H:n]] 后面 —— 否则会留下一个没有标题的空标记，
            # 前端目录会把紧随其后的正文行吞成标题。
            text = re.sub(
                r"(\[\[H:\d+(?:@\d+)?\]\])\s*" + re.escape(marker),
                lambda m: f"{m.group(1)} {raw.strip()}",
                text,
            )
            text = text.replace(marker, raw.strip())
            changed = True
            continue
        kept.append(f)
    if changed:
        formulas[:] = kept
        text = re.sub(r"\n{3,}", "\n\n", text)
    return text


_MAX_FIGURE_BYTES = 2_000_000
_MAX_FIGURES = 60

_FIGURE_CAPTION_RE = re.compile(
    r"(?i)(?:^|[\s　])(?:fig(?:ure)?\.?|table\.?|图|表)\s*\d+[^\n]{0,220}"
)
# 图注识别要比正文引用严格：Science 论文正文里的「Figure 1D shows ...」
# 不能被当成图注挂到最近的图上。真实图注通常从 Fig./Figure/Table/图/表开头，
# 且编号后不是紧跟面板字母再接普通单词。
_FIGURE_CAPTION_START_RE = re.compile(
    # 编号后紧跟逗号/分号通常是正文引用（如 ``Figure 2, B to D, shows…``），
    # 不是图注；真正图注常用句点、冒号、破折号或中文空格承接标题。
    r"(?i)^\s*(?:fig(?:ure)?\.?|table\.?|图|表)\s*\d+(?![A-Za-z])\s*(?![,;])[^\n]{0,220}"
)


def _caption_near(
    text_lines: list[tuple[str, list[float]]], image_bbox: list[float]
) -> str:
    x0, y0, x1, y1 = image_bbox
    candidates: list[tuple[float, str]] = []
    for text, tb in text_lines:
        tx0, ty0, tx1, ty1 = tb
        vertical_gap = ty0 - y1
        if -5 <= vertical_gap <= 180:
            score = abs(tx0 - x0) + max(0.0, vertical_gap)
            candidates.append((score, text))
        elif -80 <= ty0 - y0 <= 0:
            score = abs(tx0 - x0) + abs(y0 - ty0) + 15
            candidates.append((score, text))
    if candidates:
        candidates.sort(key=lambda item: item[0])
        for _, text in candidates:
            match = _FIGURE_CAPTION_START_RE.search(text)
            if match:
                return match.group(0).strip()
    # 近邻窗口里没找到图注 → 退一步按「整页最近的一条图注样文本」兜底。
    # 图注并不总贴着图（居中图注、双栏排版、表注在图上方 200pt 外都会落出窗口）。
    return _caption_nearest_on_page(text_lines, image_bbox)


# 页内兜底时允许的纵向距离上限（pt）：超过这个距离的图注多半属于另一张图
_CAPTION_PAGE_MAX_GAP = 260.0


def _caption_nearest_on_page(
    text_lines: list[tuple[str, list[float]]], image_bbox: list[float]
) -> str:
    """整页找最像图注的一行（含「图 N / Fig. N / Table N」），按纵向距离取最近。

    只认命中 _FIGURE_CAPTION_RE 的行，因此不会把期刊页眉/正文误当图注。
    """
    x0, y0, x1, y1 = image_bbox
    best: tuple[float, str] | None = None
    for text, tb in text_lines:
        match = _FIGURE_CAPTION_START_RE.search(text)
        if not match:
            continue
        tx0, ty0 = tb[0], tb[1]
        # 在图下方为正、上方为负
        gap = (ty0 - y1) if ty0 >= y1 else (ty0 - y0)
        if gap < -_CAPTION_PAGE_MAX_GAP or gap > _CAPTION_PAGE_MAX_GAP:
            continue
        score = abs(gap) + abs(tx0 - x0) * 0.5
        if best is None or score < best[0]:
            best = (score, match.group(0).strip())
    # 未命中"图/表"关键词时直接返回空，绝不能返回期刊页眉
    return best[1] if best else ""


_MIN_FIGURE_EDGE = 24.0
_MIN_VECTOR_FIGURE_EDGE = 80.0
_DRAWING_MERGE_GAP = 35.0
# 「框住正文」的判据：框内「正文长行」的条数。
# 正文的本质是成段排满的长文本行，而图表内只有坐标轴/图例/面板标签这类短行，
# 最多再加一两行图注。详见 _rect_prose_lines。
_PROSE_LINE_WEIGHT = 12.0  # 单行折算成汉字当量后，达到此长度才算「正文长行」
_MAX_FIGURE_PROSE_LINES = 4  # 框内正文长行数达到此值即判为「框住正文」
# 水印/底纹矢量层：纯描边无填充、由大量细碎路径拼成、且覆盖大半个版面。
_WATERMARK_MIN_ITEMS = 20
_WATERMARK_MIN_AREA_RATIO = 0.05
# 无图注且最长边小于此值的小图 = 刊头 logo / 页眉装饰，不是正文图表。
_DECOR_FIGURE_MAX_EDGE = 90.0
# 空白装饰图判据（无图注时生效）：整幅最暗处的亮度下限 + 灰度标准差上限。
# 两者必须同时成立；实测真图的最暗亮度 ≤124，装饰渐变带为 200。
_BLANK_FIGURE_MIN_LUMA = 170.0
_FLAT_FIGURE_MAX_STD = 36.0
# 近白装饰底纹里偶尔露出一小截红色/黑色页眉文字，会把「最暗亮度」拉低；
# 用像素占比补一道门，避免这种局部噪声让整块空白装饰漏网。
_NEAR_WHITE_FIGURE_MIN_RATIO = 0.985
_DARK_FIGURE_MAX_RATIO = 0.01
# 模板装饰带判据：宽高比下限 + 「同尺寸跨页重现」的最少页数
_BANNER_MIN_ASPECT = 4.0
_BANNER_MIN_REPEAT = 3
# 单页刊头/模板横条：无图注、贴近页顶且高度很薄。真实正文图通常有图注，
# 即便没有图注也不会同时满足「页顶 + 450pt 以上宽 + 80pt 以下高」这组条件。
_TOP_BANNER_MIN_WIDTH = 450.0
_TOP_BANNER_MAX_HEIGHT = 80.0
_TOP_BANNER_MAX_Y = 90.0
# 浏览器可直接渲染的图片格式；其余（jpx/jp2/jbig2/ccitt 等）需转码为 PNG
_BROWSER_IMAGE_EXTS = {"png", "jpg", "gif", "webp", "svg", "bmp"}


def _image_xref_to_png(doc: Any, xref: int) -> bytes | None:
    """把 PDF 内嵌图片（任意编码）经 Pixmap 转码为 PNG 字节。"""
    try:
        import pymupdf as fitz
    except Exception:
        try:
            import fitz
        except Exception:
            return None
    try:
        pix = fitz.Pixmap(doc, xref)
        if pix.colorspace is not None and pix.colorspace.n > 3:
            pix = fitz.Pixmap(fitz.csRGB, pix)
        if pix.alpha:
            pix = fitz.Pixmap(pix, 0)
        return pix.tobytes("png")
    except Exception:
        return None


def _page_text_lines(page: Any) -> list[tuple[str, list[float]]]:
    """从页面提取文本行及其 bbox，供图表标题搜索和区域排除使用。"""
    try:
        text_dict = page.get_text("dict")
    except Exception:
        return []
    lines: list[tuple[str, list[float]]] = []
    for block in text_dict.get("blocks", []):
        if block.get("type") != 0:
            continue
        for line in block.get("lines", []):
            line_text = "".join(
                str(span.get("text", "")) for span in line.get("spans", [])
            ).strip()
            bbox = line.get("bbox")
            if line_text and bbox:
                lines.append((line_text, list(bbox)))
    return lines


_CJK_RE = re.compile(r"[\u3400-\u4dbf\u4e00-\u9fff\uf900-\ufaff]")
_EN_WORD_RE = re.compile(r"[A-Za-z][A-Za-z\-']{1,}")


def _prose_weight(text: str) -> float:
    """把一行文本折算成「汉字当量」，用于跨语言比较正文体量。

    汉字计 1，英文单词计 1.5 —— 一个英文单词平均 5~6 个字符，信息量却只相当于
    1~2 个汉字。直接数字符会严重高估英文：实测英文教材里「一张曲线图 + 一行
    图注」的字符数就有 150~200，按字符数判会把真图成批误删。
    """
    s = str(text or "")
    return len(_CJK_RE.findall(s)) + 1.5 * len(_EN_WORD_RE.findall(s))


def _rect_prose_lines(
    rect: list[float], text_lines: list[tuple[str, list[float]]]
) -> int:
    """统计中心点落在 rect 内的「正文长行」条数。

    为什么这么判（两条弯路都走过）：
    - 按总字符数：对语言不公平，英文图注一行就上百字符，真图被误删（实测 THz
      教材的坐标图成批中招）；
    - 按总行数：图表内的坐标轴刻度、图例、面板标签贡献大量短行，实测热障涂层
      某图框内 32 行却只有 89 字的标注。
    正文的真正特征是「成段的、每行都排满的长文本行」（中文一行 30+ 字，英文
    一行 60+ 字符），而图表最多只有一两行图注能达到这个长度。
    """
    if not rect or len(rect) < 4:
        return 0
    try:
        x0, y0, x1, y1 = (float(v) for v in rect[:4])
    except Exception:
        return 0
    n = 0
    for text, tb in text_lines:
        if not tb or len(tb) < 4:
            continue
        try:
            cx = (float(tb[0]) + float(tb[2])) / 2.0
            cy = (float(tb[1]) + float(tb[3])) / 2.0
        except Exception:
            continue
        if x0 - 1 <= cx <= x1 + 1 and y0 - 1 <= cy <= y1 + 1:
            if _prose_weight(text) >= _PROSE_LINE_WEIGHT:
                n += 1
    return n


def _is_watermark_drawing(drawing: Any, page: Any) -> bool:
    """判断一个矢量绘图是否为水印/底纹层，而不是图表内容。

    期刊 PDF 的防盗用水印（如知网的斜纹 logo）特征相当统一：纯描边无填充、
    由大量细碎路径项拼成、且覆盖大半个版面。图表内容通常是填充式或填充与
    描边混合，路径项也远没有这么碎。

    水印必须整体跳过：它横跨正文区，一旦参与矩形合并就会把整段正文圈进
    「图表区域」，最后渲染成一张正文截图。
    """
    if not isinstance(drawing, dict):
        return False
    if drawing.get("fill") is not None:
        return False
    items = drawing.get("items") or []
    if len(items) < _WATERMARK_MIN_ITEMS:
        return False
    rect = drawing.get("rect")
    if rect is None:
        return False
    try:
        page_rect = page.rect
        area = float(rect.width) * float(rect.height)
        page_area = float(page_rect.width) * float(page_rect.height)
    except Exception:
        return False
    if page_area <= 0:
        return False
    return area / page_area >= _WATERMARK_MIN_AREA_RATIO


def _page_figures(
    doc: Any,
    page: Any,
    page_index: int,
    start_id: int,
    max_figures: int = _MAX_FIGURES,
    dropped: list[dict[str, Any]] | None = None,
) -> list[dict[str, Any]]:
    """抽取一页的全部位图（带或不带标题），转 data URI 并记录 y 坐标供原位插入。

    过滤过小图片（图标/水印/内嵌徽标），并跳过体积超限的位图，避免把整页扫描件当图。

    `dropped`：可选的收集列表。被判为「框住正文的伪图」而剔除的候选会记进这里。
    调用方据此区分「这页本来就没图」和「这页的候选全被判定为伪图」—— 后者不该
    再触发按图注兜底框取，否则剔除伪图反而会凭空多出一批新图。
    """
    try:
        image_infos = page.get_image_info(xrefs=True) or []
    except Exception:
        return []
    text_lines = _page_text_lines(page)

    figures: list[dict[str, Any]] = []
    for info in image_infos:
        bbox = info.get("bbox")
        xref = info.get("xref", 0)
        if not bbox or not xref:
            continue
        x0, y0, x1, y1 = [float(v) for v in bbox]
        if (x1 - x0) < _MIN_FIGURE_EDGE or (y1 - y0) < _MIN_FIGURE_EDGE:
            continue
        try:
            extracted = doc.extract_image(xref)
        except Exception:
            continue
        image_bytes = extracted.get("image")
        if not image_bytes or len(image_bytes) > _MAX_FIGURE_BYTES:
            continue
        # 位图同样可能整块压在正文上（页面底纹图、带文本层的扫描页）：
        # 框内出现成段的正文长行，说明它承载的是正文，收下来只会在正文里重复一遍内容。
        if _rect_prose_lines([x0, y0, x1, y1], text_lines) >= _MAX_FIGURE_PROSE_LINES:
            if dropped is not None:
                dropped.append({"bbox": [x0, y0, x1, y1]})
            continue
        caption = _caption_near(text_lines, [x0, y0, x1, y1])
        ext = str(extracted.get("ext") or "png").lower()
        if ext == "jpeg":
            ext = "jpg"
        # 浏览器不认的格式（jpx/jp2/jbig2 等）转码为 PNG，否则前端加载失败
        if ext not in _BROWSER_IMAGE_EXTS:
            png_bytes = _image_xref_to_png(doc, xref)
            if png_bytes and len(png_bytes) <= _MAX_FIGURE_BYTES:
                image_bytes = png_bytes
                ext = "png"
            else:
                # 转码失败/超限：退化为整区域渲染截图
                region = _render_figure_region(page, [x0, y0, x1, y1])
                if not region:
                    continue
                figures.append(
                    {
                        "id": f"fig{start_id + len(figures)}",
                        "page": page_index,
                        "y": y0,
                        "bbox": [x0, y0, x1, y1],
                        "caption": caption,
                        "image": region["image"],
                        "width": region["width"],
                        "height": region["height"],
                    }
                )
                if len(figures) >= max_figures:
                    break
                continue
        figures.append(
            {
                "id": f"fig{start_id + len(figures)}",
                "page": page_index,
                "y": y0,
                "bbox": [x0, y0, x1, y1],
                "caption": caption,
                "image": (
                    "data:image/"
                    + ext
                    + ";base64,"
                    + base64.b64encode(image_bytes).decode("ascii")
                ),
                "width": info.get("width"),
                "height": info.get("height"),
            }
        )
        if len(figures) >= max_figures:
            break
    return figures


def _merge_nearby_rects(
    rects: list[list[float]], gap: float = _DRAWING_MERGE_GAP
) -> list[list[float]]:
    """合并空间上相邻或重叠的矩形区域。

    按 y 坐标排序后，将水平重叠且垂直间距 < gap 的矩形合并。
    """
    if not rects:
        return []
    # 按 y0 排序，保证自顶向下合并
    sorted_rects = sorted(rects, key=lambda r: (r[1], r[0]))
    merged: list[list[float]] = [list(sorted_rects[0])]
    for rect in sorted_rects[1:]:
        x0, y0, x1, y1 = rect
        prev = merged[-1]
        # 水平方向有重叠且垂直方向足够近 → 合并
        h_overlap = x0 <= prev[2] + gap and x1 + gap >= prev[0]
        v_close = y0 <= prev[3] + gap
        if h_overlap and v_close:
            prev[0] = min(prev[0], x0)
            prev[1] = min(prev[1], y0)
            prev[2] = max(prev[2], x1)
            prev[3] = max(prev[3], y1)
        else:
            merged.append([x0, y0, x1, y1])
    return merged


def _render_figure_region(
    page: Any, bbox: list[float], dpi: int = 200, pad: float = 8.0
) -> dict[str, Any] | None:
    """从页面渲染图中裁剪指定区域，返回 {image, width, height} 或 None。"""
    try:
        import pymupdf as fitz
    except Exception:
        import fitz

    x0, y0, x1, y1 = [float(v) for v in bbox]
    x0 = max(0.0, x0 - pad)
    y0 = max(0.0, y0 - pad)
    x1 = min(float(page.rect.width), x1 + pad)
    y1 = min(float(page.rect.height), y1 + pad)
    if x1 - x0 < _MIN_FIGURE_EDGE or y1 - y0 < _MIN_FIGURE_EDGE:
        return None

    clip = fitz.Rect(x0, y0, x1, y1)
    matrix = fitz.Matrix(dpi / 72, dpi / 72)
    try:
        pix = page.get_pixmap(clip=clip, matrix=matrix, alpha=False)
    except Exception:
        return None
    img_bytes = pix.tobytes("png")
    if len(img_bytes) > _MAX_FIGURE_BYTES:
        return None

    return {
        "image": "data:image/png;base64," + base64.b64encode(img_bytes).decode("ascii"),
        "width": pix.width,
        "height": pix.height,
    }


def _rects_overlap(a: list[float], b: list[float]) -> bool:
    """两个矩形是否重叠（含边界接触）。"""
    return a[0] <= b[2] and a[2] >= b[0] and a[1] <= b[3] and a[3] >= b[1]


def _vector_figure_regions(
    page: Any,
    page_index: int,
    start_id: int,
    bitmap_bboxes: list[list[float]],
    max_figures: int = _MAX_FIGURES,
    dropped: list[dict[str, Any]] | None = None,
) -> list[dict[str, Any]]:
    """`dropped` 语义同 `_page_figures`：记录被跳过的水印层与正文框。"""
    try:
        drawings = page.get_drawings() or []
    except Exception:
        return []
    if not drawings:
        return []

    text_lines = _page_text_lines(page)
    drawing_rects: list[list[float]] = []
    simple_fill_rects: list[list[float]] = []
    for d in drawings:
        # 水印/底纹层整体跳过：它横跨正文区，一旦参与合并就会把整段正文圈进来
        if _is_watermark_drawing(d, page):
            if dropped is not None:
                rect = d.get("rect")
                if rect is not None:
                    dropped.append(
                        {
                            "bbox": [
                                float(rect.x0),
                                float(rect.y0),
                                float(rect.x1),
                                float(rect.y1),
                            ]
                        }
                    )
            continue
        rect = d.get("rect")
        if not rect:
            continue
        r = [float(rect.x0), float(rect.y0), float(rect.x1), float(rect.y1)]
        w = r[2] - r[0]
        h = r[3] - r[1]
        # 仅排除细线/边框，允许图表内部包含标注文字
        if w < 10 or h < 6:
            continue
        # 演示文稿常把标题/正文放在一个「实心填充、单个矩形路径」里。
        # 这类背景框没有任何图形内容；若框内还有两行以上文本，必然是版式容器，
        # 不能渲染成图表。真实矢量图会包含多条路径/填充色，或由位图承载主体。
        fill = d.get("fill")
        items = d.get("items") or []
        if (
            fill is not None
            and len(items) <= 2
        ):
            simple_fill_rects.append(list(r))
            # 实心底色框本身不参与图形合并；若后面有边框，会由下面的
            # containment 规则连同边框一起判成版式容器。
            continue
        if any(_rects_overlap(r, bb) for bb in bitmap_bboxes):
            continue
        drawing_rects.append(r)

    if not drawing_rects:
        return []

    merged = _merge_nearby_rects(drawing_rects, gap=40.0)
    candidate_rects = [
        r
        for r in merged
        if (r[2] - r[0]) >= _MIN_VECTOR_FIGURE_EDGE
        and (r[3] - r[1]) >= _MIN_VECTOR_FIGURE_EDGE
    ]

    # 页面钳制与垃圾矩形过滤：
    # - 合并外接矩形可能超出页面（跨页装饰线）；裁到页内
    # - 框住大段正文的"页面边框"矩形不是图表 → 剔除
    # - 图内混入的正文段落（图表与下方正文被装饰线桥接）→ 垂直收紧，再渲染
    try:
        page_rect = page.rect
        pw, ph = float(page_rect.width), float(page_rect.height)
    except Exception:
        pw = ph = 0.0
    filtered: list[list[float]] = []
    for rect in candidate_rects:
        r = list(rect)
        if pw > 0:
            r = [max(0.0, r[0]), max(0.0, r[1]), min(pw, r[2]), min(ph, r[3])]
        if r[2] - r[0] < 10 or r[3] - r[1] < 10:
            continue
        if pw > 0 and (r[2] - r[0]) * (r[3] - r[1]) > 0.85 * pw * ph:
            continue  # 近整页矩形：页面边框/装饰
        # 边框与填充通常是两个 drawing；合并后不能只看当前 drawing 的 fill。
        # 若候选完整包住实心底色框且框内有 >=2 行正文长行，它是幻灯片
        # 的文本容器（例如提问框），不是图表。
        if any(
            sr[0] >= r[0] - 2 and sr[1] >= r[1] - 2
            and sr[2] <= r[2] + 2 and sr[3] <= r[3] + 2
            and _rect_prose_lines(sr, text_lines) >= 2
            for sr in simple_fill_rects
        ):
            if dropped is not None:
                dropped.append({"bbox": list(r)})
            continue
        # 框内有成段的正文长行 → 这块区域装的是正文，不是图表。
        # 中文文献里这一条是主要兜底：下面那句英文单词判据对中文正文恒为空集。
        if _rect_prose_lines(r, text_lines) >= _MAX_FIGURE_PROSE_LINES:
            if dropped is not None:
                dropped.append({"bbox": list(r)})
            continue
        prose_bboxes = [
            tb
            for txt, tb in text_lines
            if len(re.findall(r"[A-Za-z]{2,}", str(txt))) >= 5
            and tb[0] >= r[0] - 2
            and tb[2] <= r[2] + 2
            and tb[1] >= r[1] - 2
            and tb[3] <= r[3] + 2
        ]
        if len(prose_bboxes) > 6:
            continue  # 框住大段正文 → 装饰框
        # 垂直收紧：图下方（或上方）被桥接进来的正文段落，从矩形里裁出去
        rh = r[3] - r[1]
        for tb in sorted(prose_bboxes, key=lambda b: b[1]):
            if tb[1] > r[1] + 0.25 * rh and tb[3] < r[3] - 6:
                r[3] = tb[1] - 4  # 下方正文 → 截断底部
                break
        rh = r[3] - r[1]
        for tb in sorted(prose_bboxes, key=lambda b: -b[3]):
            if tb[3] < r[3] - 0.25 * rh and tb[1] > r[1] + 6:
                r[1] = tb[3] + 4  # 上方正文 → 截断顶部
                break
        if r[2] - r[0] < _MIN_VECTOR_FIGURE_EDGE or r[3] - r[1] < _MIN_VECTOR_FIGURE_EDGE:
            continue
        filtered.append(r)
    candidate_rects = filtered

    figures: list[dict[str, Any]] = []
    for rect in candidate_rects:
        if len(figures) >= max_figures:
            break
        rendered = _render_figure_region(page, rect, dpi=200)
        if rendered is None:
            continue
        caption = _caption_near(text_lines, rect)
        figures.append(
            {
                "id": f"fig{start_id + len(figures)}",
                "page": page_index,
                "y": rect[1],
                "bbox": rect,
                "caption": caption,
                "image": rendered["image"],
                "width": rendered.get("width"),
                "height": rendered.get("height"),
                "source": "vector",
            }
        )
    return figures


def _caption_driven_figure_fallback(
    doc: Any, page: Any, page_index: int, start_id: int
) -> list[dict[str, Any]]:
    """以页面中已明确检测到的 '图 N' 标题为基准，向上框取渲染完整图表区域"""
    text_lines = _page_text_lines(page)
    caption_candidates = []
    for text, bbox in text_lines:
        m = _FIGURE_CAPTION_START_RE.search(text)
        if m and ("图" in text or "Fig" in text):
            caption_candidates.append((text.strip(), bbox))

    fallback_figs = []
    for caption_text, cb in caption_candidates:
        cx0, cy0, cx1, cy1 = cb
        top_y = max(60.0, cy0 - 320.0)
        for txt, tb in text_lines:
            if tb[3] <= cy0 - 20 and tb[3] > top_y:
                if not _FIGURE_CAPTION_START_RE.search(txt) and len(txt) > 20:
                    top_y = max(top_y, tb[3] + 4)

        fig_bbox = [
            max(40.0, cx0 - 80),
            top_y,
            min(float(page.rect.width) - 40, cx1 + 80),
            cy0 - 4,
        ]
        # 框内出现成段正文 → 圈住的是图注上方的正文，不是图（判据同 _rect_prose_lines）
        if _rect_prose_lines(fig_bbox, text_lines) >= _MAX_FIGURE_PROSE_LINES:
            continue
        if (fig_bbox[2] - fig_bbox[0] >= 100) and (fig_bbox[3] - fig_bbox[1] >= 80):
            rendered = _render_figure_region(page, fig_bbox, dpi=200)
            if rendered:
                fallback_figs.append(
                    {
                        "id": f"fig{start_id + len(fallback_figs)}",
                        "page": page_index,
                        "y": fig_bbox[1],
                        "bbox": fig_bbox,
                        "caption": caption_text,
                        "image": rendered["image"],
                        "width": rendered["width"],
                        "height": rendered["height"],
                        "source": "caption_crop",
                    }
                )
    return fallback_figs


# 面板标签位于面板左上角，标签上沿距面板内容上沿通常只有几个 pt，这里留足上边距以包含顶层面板内容
_PANEL_TOP_PAD = 18.0


def _caption_panel_letters(caption_text: str) -> set[str]:
    """从文本中的括号子图引用（如 (A)、(E and F)）提取候选面板字母。"""
    if not caption_text:
        return set()
    letters: set[str] = set()
    for m in re.finditer(r"\(\s*([A-Za-z])\s*(?:and\s*([A-Za-z])\s*)?\)", caption_text):
        letters.add(m.group(1).upper())
        if m.group(2):
            letters.add(m.group(2).upper())
    return letters


def _detect_panel_labels(
    text_lines: list[tuple[str, list[float]]],
    top_y: float,
    cap_y: float,
    caption_text: str = "",
) -> list[dict[str, Any]]:
    """在图区域内扫描独立的大写字母（A、B、(C)）作为子图标签。

    双保险确认真正的面板标签：
    (1) 文本行本身是单个字母（裸 A/B/C…，轴上的 X/Y/µm 也是字母，需配合第 2 步过滤）；
    (2) 其字母出现在图注的括号子图引用里（(A)...(H)）。
    因为版面图注常分左右两栏，右栏的 (C)~(H) 与左栏起始行不同行，故括号引用须取自整页文本而非仅图注首行。
    """
    expected: set[str] = set()
    for _t, _tb in text_lines:
        expected |= _caption_panel_letters(str(_t))
    singles: list[dict[str, Any]] = []
    # 裸字母仅接受大写（坐标轴 x/y/z 是小写，不能当面板标签）；
    # 书籍/期刊常用小写括号标签 (a)(b) → 括号形式大小写均可
    panel_letter_re = re.compile(r"^\s*(?:\(\s*([A-Za-z])\s*\)|([A-Z]))\s*$")
    for txt, tb in text_lines:
        if not tb or not (top_y - 4 <= tb[1] <= cap_y):
            continue
        m = panel_letter_re.match(str(txt))
        if not m:
            continue
        letter = (m.group(1) or m.group(2) or "").strip().upper()
        if not letter:
            continue
        singles.append(
            {
                "label": letter,
                "x0": tb[0], "y0": tb[1], "x1": tb[2], "y1": tb[3],
            }
        )
    if expected and singles:
        kept = [s for s in singles if s["label"] in expected]
        # 仅当按图注过滤后仍能构成完整的子图标注序列时，才采信过滤结果
        if len(kept) >= 2:
            singles = kept
    elif not expected and singles:
        # 图注无括号引用时：仅当字母构成连续序列（A,B,…）才采信，
        # 避免把图内孤立坐标轴字母（X/Y/Z）误判成面板
        if not {s["label"] for s in singles} <= set("ABCDEFGH"[: len(singles)]):
            singles = []
    # 同一字母可能命中面板标签与面板内坐标/图例旁的孤立字母，只保留最靠上靠左的真实标签
    _best: dict[str, dict[str, Any]] = {}
    for p in singles:
        cur = _best.get(p["label"])
        if cur is None or (p["y0"], p["x0"]) < (cur["y0"], cur["x0"]):
            _best[p["label"]] = p
    singles = list(_best.values())
    # 按阅读顺序（先上后下、先左后右）排序，保证 A/B/C…与版面一致；
    # y 加 10pt 行容差：同一行的 (a)(b) 标签 y 常差零点几 pt，纯按 y 排会颠倒左右
    singles.sort(key=lambda p: (round(p["y0"] / 10.0), p["x0"]))
    return singles


def _fig_caption_number(caption_text: str) -> str:
    m = re.search(r"(?i)(?:fig(?:ure)?\.?|table\.?|图|表)\s*(\d+)", caption_text or "")
    return m.group(1) if m else ""


def _layout_figure_regions(
    page: Any, page_figs: list[dict[str, Any]]
) -> list[dict[str, Any]]:
    """按版面把同页同一幅多子图合并为整幅区域，恢复被拆散的面板并为每块挂上整幅引用。

    区域边界策略（不再依赖易受正文/水印干扰的文本行启发式）：
    - 垂直带：上一条图注底边 → 本条图注顶边，直接把带内所有碎片都视为本图成员，
      （否则空标题的子面板碎片（如单独占据的 A）会被漏掉，导致「只截取一部分」）。
    - 上边界：由最上层子图标签上沿外推（标签总在面板内容上缘附近），避免正文长行/「Downloaded」等水印被误判。
    - 水平边界：成员碎片与子图标签外接并扩充，夹到页边距内。
    返回 [{ bbox, figure_number, panels:[{label,x0,y0,x1,y1}], members:[fig,...] }]
    """
    try:
        mw = float(page.rect.width)
    except Exception:
        mw = 594.0
    text_lines = _page_text_lines(page)
    # 识别图注：从“图 N/表 N”起始行向下连续吸收续行，拼成完整图注文本与整块 bbox
    captions: list[tuple[str, list[float], list[float]]] = []
    idx = 0
    n = len(text_lines)
    while idx < n:
        txt, tb = text_lines[idx]
        t = str(txt)
        if tb and _FIGURE_CAPTION_START_RE.search(t) and (
            "fig" in t.lower() or "table" in t.lower() or "图" in t or "表" in t
        ):
            text = t.strip()
            y0, yb = float(tb[1]), float(tb[3])
            x0 = float(tb[0])
            span_min, span_max = y0, yb
            j = idx + 1
            while j < n:
                ntxt, ntb = text_lines[j]
                if not ntb or float(ntb[1]) - yb > 12 or abs(float(ntb[0]) - x0) > 30:
                    break
                text += " " + str(ntxt).strip()
                yb = float(ntb[3])
                span_min = min(span_min, float(ntb[1]))
                span_max = max(span_max, yb)
                j += 1
            captions.append((text, [x0, span_min, float(tb[2]), span_max], [x0, span_min, float(tb[2]), span_max]))
            idx = j
        else:
            idx += 1
    if not captions:
        return []
    captions.sort(key=lambda c: c[1][1])

    def _center_y(fig: dict[str, Any]) -> float | None:
        b = fig.get("bbox")
        if not b or len(b) < 4:
            return None
        return (b[1] + b[3]) / 2.0

    regions: list[dict[str, Any]] = []
    band_top = 8.0
    for caption_text, cb, bb in captions:
        band_bot = float(bb[1])
        members = [
            f for f in page_figs
            if (cy := _center_y(f)) is not None and band_top - 6 <= cy <= band_bot
        ]
        if not members:
            band_top = band_bot
            continue

        xs = [float(f["bbox"][0]) for f in members if f.get("bbox")]
        xe = [float(f["bbox"][2]) for f in members if f.get("bbox")]
        ys = [float(f["bbox"][1]) for f in members if f.get("bbox")]
        if not xs or not ys:
            band_top = band_bot
            continue
        x0, x1 = min(xs), max(xe)

        panels = _detect_panel_labels(text_lines, band_top - 4, band_bot, caption_text)
        if panels:
            x0 = min(x0, min(p["x0"] for p in panels))
            x1 = max(x1, max(p["x1"] for p in panels))
            top = min(p["y0"] for p in panels) - _PANEL_TOP_PAD
        else:
            top = min(ys) - 6.0
        # 上边界夹到「不高于页顶/前一幅图注」的合理区间，防止过截到正文顶部
        top = max(top, band_top - 8, 22.0)
        # 下边界取成员内容（及面板标签）下沿，留少量边距即可，避免图注上方的大片空白被吞进裁剪
        content_bottom = max(max(float(f["bbox"][3]) for f in members if f.get("bbox")),
                             max((p["y1"] for p in panels), default=0.0)) + 6.0
        y1 = max(top + 40, min(band_bot - 3, content_bottom))
        x0 = max(30.0, x0 - 10)
        x1 = min(mw - 30, x1 + 10)
        # 正文闯入收缩：成员矩形偶发把上下方正文段落桥接进区域，
        # 与区域上/下边缘相交的正文长行（≥5 个词、横向占区域过半）把边界推回去
        for txt, tb in text_lines:
            s = str(txt)
            if not tb or len(re.findall(r"[A-Za-z]{2,}", s)) < 5:
                continue
            lx0, ly0, lx1, ly1 = (float(v) for v in tb)
            if min(x1, lx1) - max(x0, lx0) < 0.5 * (lx1 - lx0):
                continue
            if top <= ly1 <= top + 22:
                top = ly1 + 2
            if y1 - 22 <= ly0 <= y1:
                y1 = ly0 - 2
        if (x1 - x0) < 60 or (y1 - top) < 40:
            band_top = band_bot
            continue
        regions.append(
            {
                "bbox": [x0, top, x1, y1],
                "figure_number": _fig_caption_number(caption_text),
                "panels": panels,
                "members": members,
            }
        )
        band_top = band_bot
    return regions


def _attach_whole_ref(
    page_figs: list[dict[str, Any]], regions: list[dict[str, Any]]
) -> None:
    """把整幅区域引用（whole_ref）挂到属于它的每个子块上，供前端合并渲染。"""
    for region in regions:
        ref = {
            "bbox": region["bbox"],
            "figure_number": region["figure_number"],
            "panels": region["panels"],
        }
        for fig in region["members"]:
            fig["whole_ref"] = ref


def _is_blank_figure(figure: dict[str, Any]) -> bool:
    """位图本身「无墨迹且几乎没明暗起伏」= 底纹、渐变装饰条，不是图表。

    实测判据（同一份幻灯片 PDF：9 张装饰带 vs 7 张真图完全可分）：
      装饰带 最暗亮度 200、灰度标准差 11；真图 最暗亮度 ≤124、标准差 ≥52。
    因此要求「两者同时成立」才判装饰 —— 只判其一都会误伤浅色图（如只有细轴线、
    最暗亮度 124、标准差 52 的曲线图）。无图注是前置条件，见 _is_decor_figure。
    """
    data_uri = str(figure.get("image") or "")
    if not data_uri.startswith("data:"):
        return False
    try:
        raw = base64.b64decode(data_uri.split(",", 1)[1])
    except Exception:
        return False
    if not raw:
        return True
    try:
        from PIL import Image, ImageStat
    except Exception:
        return False  # 没装 Pillow 就跳过这条判据，绝不因此漏掉图
    try:
        with Image.open(BytesIO(raw)) as img:
            gray = img.convert("L")
            extrema = gray.getextrema()
            spread = float(ImageStat.Stat(gray).stddev[0])
            pixels = list(gray.resize((min(160, gray.width), min(160, gray.height))).getdata())
        darkest = float(extrema[0])
    except Exception:
        return False
    near_white = sum(v >= 220 for v in pixels) / max(1, len(pixels))
    dark_ratio = sum(v < _BLANK_FIGURE_MIN_LUMA for v in pixels) / max(1, len(pixels))
    return (
        (darkest >= _BLANK_FIGURE_MIN_LUMA and spread <= _FLAT_FIGURE_MAX_STD)
        or (
            near_white >= _NEAR_WHITE_FIGURE_MIN_RATIO
            and dark_ratio <= _DARK_FIGURE_MAX_RATIO
        )
    )


def _is_decor_figure(figure: dict[str, Any]) -> bool:
    """无图注、最长边过小的小图 = 刊头 logo / 页眉装饰，不是正文图表。

    判据只放在最终收口处，而不是各提取分支里：页眉 logo 与多子图的面板碎片
    外形上无法区分，但面板碎片要被 _layout_figure_regions 用来合成整幅图，
    提前丢掉会毁掉整幅裁剪 —— 所以只拦「已经不会再被谁引用」的成品。
    """
    if str(figure.get("caption") or "").strip():
        return False
    if _is_blank_figure(figure):
        return True
    bbox = figure.get("bbox")
    if not bbox or len(bbox) < 4:
        return False
    try:
        x0, y0, x1, y1 = (float(v) for v in bbox[:4])
    except Exception:
        return False
    # 期刊首页常把刊名/栏目做成一条整幅彩色图片。它不是「近白空图」，
    # 也只出现一页，因此重复横条规则覆盖不到；版面上可由贴顶、超宽、超薄
    # 且无图注这组稳定信号识别。
    if (
        y0 <= _TOP_BANNER_MAX_Y
        and (x1 - x0) >= _TOP_BANNER_MIN_WIDTH
        and (y1 - y0) <= _TOP_BANNER_MAX_HEIGHT
    ):
        return True
    return max(abs(x1 - x0), abs(y1 - y0)) < _DECOR_FIGURE_MAX_EDGE


def _drop_repeated_banner_figures(figures: list[dict[str, Any]]) -> set[str]:
    """无图注 + 同一尺寸在多页重复出现的扁长条 = 模板装饰带（页眉色块/渐变条）。

    单张装饰带与真图难分，但「同一尺寸在 ≥3 页上原样重现」是模板的强特征：
    真图表极少跨页保持像素级同尺寸。再加宽高比 ≥4 的门槛，错杀面进一步收窄。
    实测：某幻灯片 PDF 的 16 张「图」里有 9 张是这条 588×96 的渐变带。
    """
    groups: dict[tuple[int, int], list[dict[str, Any]]] = {}
    for figure in figures:
        if str(figure.get("caption") or "").strip():
            continue
        bbox = figure.get("bbox")
        if not bbox or len(bbox) < 4:
            continue
        try:
            width = int(round(float(bbox[2]) - float(bbox[0])))
            height = int(round(float(bbox[3]) - float(bbox[1])))
        except Exception:
            continue
        if width <= 0 or height <= 0 or width / height < _BANNER_MIN_ASPECT:
            continue
        groups.setdefault((width, height), []).append(figure)

    dropped: set[str] = set()
    for members in groups.values():
        pages = {m.get("page") for m in members}
        if len(members) < _BANNER_MIN_REPEAT or len(pages) < _BANNER_MIN_REPEAT:
            continue
        dropped.update(str(m.get("id")) for m in members if m.get("id"))
    return dropped


def _finalize_figures_and_markers(
    figures: list[dict[str, Any]], page_texts: list[str]
) -> tuple[list[dict[str, Any]], list[str]]:
    """过滤期刊元数据图、刊头小图、空白装饰图与模板装饰带，重排 fig id 并同步重写标记。"""
    dropped_ids: set[str] = _drop_repeated_banner_figures(figures)
    kept: list[dict[str, Any]] = []
    for f in figures:
        fid = str(f.get("id") or "")
        if (
            fid in dropped_ids
            or _looks_like_journal_metadata(f.get("caption", ""))
            or _is_decor_figure(f)
        ):
            if fid:
                dropped_ids.add(fid)
            continue
        kept.append(f)
    figures = kept
    # 被剔除的图，正文里的 [[FIG:id]] 标记必须一起清掉：
    # 悬空标记会被前端当成「未能提取的图片」渲染成一个占位块。
    if dropped_ids:
        drop_re = re.compile(
            r"\[\[FIG:(?:"
            + "|".join(re.escape(i) for i in sorted(dropped_ids))
            + r")\]\]"
        )
        page_texts = [
            re.sub(r"\n{3,}", "\n\n", drop_re.sub("", pt)) for pt in page_texts
        ]
    old_to_new: dict[str, str] = {}
    for i, f in enumerate(figures):
        old_id = str(f.get("id") or "")
        new_id = f"fig{i}"
        f["id"] = new_id
        if old_id and old_id != new_id:
            old_to_new[old_id] = new_id
    if old_to_new:
        marker_re = re.compile(r"\[\[FIG:([^\]]+)\]\]")

        def _sub(match: re.Match[str]) -> str:
            old = match.group(1)
            return f"[[FIG:{old_to_new.get(old, old)}]]"

        page_texts = [marker_re.sub(_sub, pt) for pt in page_texts]
    return figures, page_texts


_FIG_MARKER_RE = re.compile(r"\[\[FIG:([^\]]+)\]\]")
_HEADING_MARKER_LINE_RE = re.compile(r"^\[\[H:\d+(?:@\d+)?\]\][ \t]*(.*)$")


def attach_figure_contexts(text: str, figures: list[dict[str, Any]]) -> None:
    """给「没有图注」的图挂上 context = 它上面最近的那条章节标题。

    只补上下文，**不编造图号** —— 正文里没有「图 N」时硬派一个号码，会让用户以为
    原文真有这个编号。前端图表列表据此显示「所属小节 + 页码」，比干巴巴的
    「第 N 页图表」有用得多，也仍然诚实。
    """
    if not figures or "[[FIG:" not in text:
        return
    by_id = {str(f.get("id")): f for f in figures if f.get("id")}
    if not by_id:
        return
    current = ""
    for raw_line in text.split("\n"):
        line = raw_line.strip()
        if not line:
            continue
        head = _HEADING_MARKER_LINE_RE.match(line)
        if head:
            current = head.group(1).strip()
            continue
        for fig_id in _FIG_MARKER_RE.findall(line):
            figure = by_id.get(fig_id)
            if figure is None:
                continue
            if str(figure.get("caption") or "").strip():
                continue
            if str(figure.get("context") or "").strip():
                continue
            if current:
                figure["context"] = current


_PIPE_ARTIFACT_RE = re.compile(r"\|+(?:[^\S\n]+\|+)+")
# 行首/行尾的 2+ 连续竖线（PDF 文本层里跨行大括号的残片，如 `||| -∫ 0 z μt(u) du`）。
# 要求竖线后紧跟空白或行尾，避免误伤行首的 ||x|| 范数写法；单竖线可能是绝对值，不动。
_PIPE_LINE_EDGE_RE = re.compile(
    r"(?m)^[ \t]*\|{2,}[ \t]+|^[ \t]*\|{2,}[ \t]*$|(?<=[ \t])\|{2,}[ \t]*$"
)


def _strip_pipe_artifacts(text: str) -> str:
    """清除 PDF 文本层里公式/表格排版残留的竖线串（如 `| | F 0( ) κ = 0 | | |`、行首 `|||`）。

    仅匹配「竖线之间只隔空白」的串与行首/行尾的 2+ 竖线簇，
    不会误伤 ||x|| 范数或 |F(x)| 绝对值写法。
    """
    text = _PIPE_ARTIFACT_RE.sub(" ", text)
    text = _PIPE_LINE_EDGE_RE.sub("", text)
    return re.sub(r"\n{3,}", "\n\n", text)


def _insert_figure_markers(
    text: str, page: Any, page_figs: list[dict[str, Any]]
) -> str:
    """按 y 坐标把 [[FIG:id]] 标记插入到正文对应位置后（PDF 自上而下阅读顺序）。"""
    if not page_figs:
        return text
    try:
        raw = page.get_text("dict")
    except Exception:
        raw = {}
    bounds: list[float] = []
    for block in raw.get("blocks", []):
        if block.get("type") != 0:
            continue
        for line in block.get("lines", []):
            bbox = line.get("bbox")
            if bbox:
                bounds.append(float(bbox[3]))  # 文本行底部 y

    lines = text.split("\n")
    inserts: list[tuple[int, str]] = []
    for fig in sorted(page_figs, key=lambda f: float(f.get("y", 0) or 0)):
        top = float(fig.get("y", 0) or 0)
        after = 0
        for i, bottom in enumerate(bounds):
            if bottom <= top:
                after = i + 1
        after = min(after, len(lines))
        inserts.append((after, f"[[FIG:{fig['id']}]]"))

    for after, marker in sorted(inserts, key=lambda item: -item[0]):
        lines.insert(after, f"\n\n{marker}\n\n")
    return "\n".join(lines)


# 行尾公式编号（如 ``ext) ⋅z] （5）`` 尾部的 ``（5）``）：
# 分类探测时先剥离——全角括号会把整行误判成正文，导致区域右侧截断
_EQ_NUM_TAIL_RE = re.compile(r"[\(（]\s*(\d{1,3}[a-zA-Z]?)\s*[\)）]\s*$")

# 分类探测用的全角标点翻译：分段函数等碎片常粘全角逗号/括号
# （如 ``1， 0， if z≥z0``），不翻译会被误判成正文——整行进不了区域，
# 还会落入正文带把相邻真碎片剪掉
_FULLWIDTH_PROBE_TRANS = str.maketrans("（）［］，．：；", "()[],.:;")


def _display_math_regions(
    page: Any, figure_bboxes: list[list[float]] | None = None
) -> list[dict[str, Any]]:
    """检测页面上由多行碎片组成的显示公式区域（文本层线性化会把它们毁掉）。

    判定思路：公式碎片在 PDF 文本层里常被拆成大量独立 block（阅读顺序不可靠），
    因此按几何位置做连通分量聚类——两碎片垂直区间重叠或间距很小、且水平重叠
    或间距 ≤30pt 时视为同一公式的零件；连通分量须 ≥2 个碎片且至少含一条硬数学行
    （含 = ∫ ∑ 等强信号）。正文行（含汉字或英文句子）天然被排除。
    返回 [{"bbox": [...], "lines": [str]}]，供区域渲染成图片。
    """
    try:
        text_dict = page.get_text("dict")
        page_w = float(page.rect.width)
        page_h = float(page.rect.height)
    except Exception:
        return []
    all_lines: list[dict[str, Any]] = []
    body_bands: list[list[float]] = []
    frags: list[dict[str, Any]] = []
    for block in text_dict.get("blocks", []):
        if block.get("type") != 0:
            continue
        for line in block.get("lines", []):
            spans = line.get("spans", [])
            text = "".join(str(s.get("text", "")) for s in spans).strip()
            bbox = line.get("bbox")
            if not text or not bbox:
                continue
            bb = [float(v) for v in bbox]
            entry: dict[str, Any] = {"text": text, "bbox": bb}
            all_lines.append(entry)
            probe = _EQ_NUM_TAIL_RE.sub("", text).strip().translate(_FULLWIDTH_PROBE_TRANS)
            hard = bool(probe) and _is_hard_math_line(probe)
            if not hard and not (probe and _is_soft_math_fragment(probe)):
                if _is_body_text_line(text):
                    body_bands.append(bb)
                continue
            entry["hard"] = hard
            frags.append(entry)
    # 预过滤：与正文行实质重叠的碎片是「句子的零件」（行内变量/行内堆叠分数），
    # 不参与区域聚类——否则它们会横跨正文行，把上下文两个独立公式桥接成一团
    def overlaps_body(fb: list[float]) -> bool:
        fh = max(4.0, fb[3] - fb[1])
        return any(
            min(fb[3], bb[3]) - max(fb[1], bb[1]) > 0.3 * fh
            and min(fb[2], bb[2]) - max(fb[0], bb[0]) > 0
            for bb in body_bands
        )

    frags = [f for f in frags if not overlaps_body(f["bbox"])]
    if len(frags) < 2:
        return []

    # 双栏版面才按栏隔离：单栏论文的宽公式可能横跨页面中线，分栏会把它们劈开
    two_col = _is_two_column_layout(all_lines, page_w)

    def col(b: list[float]) -> int:
        return 0 if (b[0] + b[2]) / 2 < page_w / 2 else 1

    n = len(frags)
    parent = list(range(n))

    def find(i: int) -> int:
        while parent[i] != i:
            parent[i] = parent[parent[i]]
            i = parent[i]
        return i

    for i in range(n):
        bi = frags[i]["bbox"]
        hi = max(4.0, bi[3] - bi[1])
        for j in range(i + 1, n):
            bj = frags[j]["bbox"]
            if two_col and col(bi) != col(bj):
                continue
            hj = max(4.0, bj[3] - bj[1])
            vgap = max(bi[1] - bj[3], bj[1] - bi[3], 0.0)
            # 阈值以正文行高为锚（大括号/积分号可高达 40+pt，若不封顶，
            # 高字形会把阈值撑到能跨过大段正文）
            if vgap > max(6.0, 1.4 * min(min(hi, hj), 18.0)):
                continue
            # 同一行（垂直重叠）的碎片允许更大水平间距：LaTeX 书排版里
            # "w2 = w1 ·  [分数主体]  1]1/2 ," 会被拆成多个远距碎片
            hgap_limit = 90.0 if vgap <= 2.0 else 30.0
            hgap = max(bi[0] - bj[2], bj[0] - bi[2], 0.0)
            if hgap > hgap_limit:
                continue
            ri, rj = find(i), find(j)
            if ri != rj:
                parent[ri] = rj

    groups: dict[int, list[int]] = {}
    for i in range(n):
        groups.setdefault(find(i), []).append(i)
    # 独立成行的公式编号（N），供区域渲染时一并纳入图片
    num_lines = [
        ln
        for ln in all_lines
        if re.fullmatch(r"[\(（]\s*(\d{1,3}[a-zA-Z]?)\s*[\)）]", ln["text"].strip())
    ]
    # 非碎片行（正文/未分类）与页面图片：供区域渲染夹取时避开邻近内容
    foreign_lines = [ln for ln in all_lines if "hard" not in ln]
    image_rects = _page_image_rects(page)
    regions: list[dict[str, Any]] = []
    for members in groups.values():
        # 一个堆叠分数就能拆出几十个碎片（每个符号/上下标都是独立碎片），
        # 成员上限要足够宽松；真正的防粘连靠高度上限
        if len(members) < 2 or len(members) > 80:
            continue
        if not any(frags[m]["hard"] for m in members):
            continue
        members.sort(key=lambda m: (frags[m]["bbox"][1], frags[m]["bbox"][0]))
        x0 = min(frags[m]["bbox"][0] for m in members)
        y0 = min(frags[m]["bbox"][1] for m in members)
        x1 = max(frags[m]["bbox"][2] for m in members)
        y1 = max(frags[m]["bbox"][3] for m in members)
        if y1 - y0 > 0.3 * page_h:  # 防整栏粘连
            continue
        lines = [frags[m]["text"] for m in members]
        # 图内标签簇（坐标轴刻度/图例，如 "z1 = 0.0"+"0.1"+"f"）：
        # 所有行都极短且区域很窄 → 不是公式，避免图表插图被当公式截图
        if (x1 - x0) < 60 and all(len(t.strip()) <= 12 for t in lines):
            continue
        # 含 (a)/(b)/(A) 面板标签行的碎片簇 → 图表内部（子图标签只出现在插图里）。
        # 用严格版正则：裸字母仅大写——公式里的单字母变量（f、z）不能触发
        if any(
            re.match(r"^\s*(?:\(\s*[A-Za-z]\s*\)|[A-Z])\s*$", t.strip()) for t in lines
        ):
            continue
        # 与页面图区域（位图/矢量图）中心相交的碎片簇 → 图内文字，不是正文公式
        if figure_bboxes:
            cx, cy = (x0 + x1) / 2, (y0 + y1) / 2
            if any(
                len(fb) == 4 and fb[0] - 2 <= cx <= fb[2] + 2 and fb[1] - 2 <= cy <= fb[3] + 2
                for fb in figure_bboxes
            ):
                continue
        # 公式右侧的编号行纳入区域：编号随图片渲染，位置保真，正文不再出现游离编号
        numbers: list[str] = []
        for ln in num_lines:
            nb = ln["bbox"]
            yc = (nb[1] + nb[3]) / 2
            # 编号常与公式基线对齐，略低于碎片区域底部 → y 窗口向下放宽一行；
            # 书籍右对齐编号距公式主体可达 100pt+ → 阈值放宽到 130
            if y0 - 4 <= yc <= y1 + 16 and nb[0] >= x1 - 20 and nb[0] - x1 <= 130:
                x1 = max(x1, nb[2])
                y0 = min(y0, nb[1])
                y1 = max(y1, nb[3])
                numbers.append(ln["text"].strip())
        # 碎片行尾粘着的编号（如 ``ext) ⋅z] （5）``）：编号已在区域图片里，
        # 记入 numbers 供调用方抑制前端角标、保留编号元数据
        for m in members:
            tm = _EQ_NUM_TAIL_RE.search(frags[m]["text"])
            if tm:
                numbers.append(tm.group(0))
        regions.append(
            {
                "bbox": [x0, y0, x1, y1],
                "lines": lines,
                "numbers": numbers,
                "clip": _clamp_math_region_clip(
                    x0, y0, x1, y1, foreign_lines, image_rects
                ),
            }
        )
    regions.sort(key=lambda r: (r["bbox"][1], r["bbox"][0]))
    return regions


def _page_image_rects(page: Any) -> list[list[float]]:
    rects: list[list[float]] = []
    try:
        for img in page.get_images(full=True):
            for r in page.get_image_rects(img[0]):
                rects.append([float(r.x0), float(r.y0), float(r.x1), float(r.y1)])
    except Exception:
        pass
    return rects


def _clamp_math_region_clip(
    x0: float,
    y0: float,
    x1: float,
    y1: float,
    foreign_lines: list[dict[str, Any]],
    image_rects: list[list[float]],
    pad: float = 2.0,
) -> list[float]:
    """计算公式区域的渲染夹取范围（含 pad）。

    大括号等高碎片的行框常溢出到相邻正文行或页眉图片里，直接把 bbox+pad 渲染
    会把页眉色条、下一行正文的残影带进公式图片。这里把伸进来的薄条（正文 ≤6pt、
    图片 ≤10pt）削掉；只夹取垂直方向，水平方向没有观测到同类问题。
    """
    clip = [x0 - pad, y0 - pad, x1 + pad, y1 + pad]
    for fb in (ln["bbox"] for ln in foreign_lines):
        if min(clip[2], fb[2]) - max(clip[0], fb[0]) <= 0:
            continue
        # 中文期刊行距紧，行底常略微落在区域顶部/底部之下，容差放宽到 ±10pt
        if clip[1] - 1.0 <= fb[3] <= y0 + 10.0:
            clip[1] = max(clip[1], fb[3] + 0.3)
        if y1 - 10.0 <= fb[1] <= clip[3] + 1.0:
            clip[3] = min(clip[3], fb[1] - 0.3)
    for rect in image_rects:
        if min(clip[2], rect[2]) - max(clip[0], rect[0]) <= 0:
            continue
        # 图片主体须在区域外（页眉/页脚带），防止削到区域内部的内容图
        if clip[1] - 0.5 <= rect[3] <= y0 + 10.0 and rect[1] < y0:
            clip[1] = max(clip[1], rect[3] + 0.3)
        if y1 - 10.0 <= rect[1] <= clip[3] + 0.5 and rect[3] > y1:
            clip[3] = min(clip[3], rect[1] - 0.3)
    if clip[3] - clip[1] < 10.0:  # 过度夹取时退回 padded bbox
        clip = [x0 - pad, y0 - pad, x1 + pad, y1 + pad]
    return clip


def _is_body_text_line(text: str) -> bool:
    """正文行：含汉字/全角标点，或含 ≥2 个英文字母，或长度 ≥8 的自然文本。"""
    s = text.strip()
    if not s:
        return False
    if re.search(r"[一-鿿　-〿＀-￯]", s):
        return True
    # 公式碎片守卫：无 ≥3 连字母词、含数字与运算符的行（如 "[(d1 −f)2 + z2"）不是正文。
    # 若被当作正文带，overlaps_body 会把 y 重叠的真公式碎片整批误杀 → 公式漏检
    if (
        len(re.findall(r"[A-Za-z]{3,}", s)) == 0
        and re.search(r"[0-9]", s)
        and re.search(r"[=+\-−–^*/()\[\]|·⋅]", s)
    ):
        return False
    if len(re.findall(r"[A-Za-z]", s)) >= 2:
        return True
    return len(s) >= 8


def _is_two_column_layout(lines: list[dict[str, Any]], page_w: float) -> bool:
    """按文本行 x 分布判断是否双栏：左右半页各有足够行、跨中线的行很少。"""
    left = right = cross = 0
    mid0, mid1 = page_w * 0.45, page_w * 0.55
    for ln in lines:
        x0, _, x1, _ = ln["bbox"]
        if x1 - x0 < 12:  # 小碎片不参与判定
            continue
        if x0 < mid0 and x1 > mid1:
            cross += 1
        elif (x0 + x1) / 2 < page_w / 2:
            left += 1
        else:
            right += 1
    total = left + right + cross
    return total >= 12 and left >= 4 and right >= 4 and cross <= 0.2 * total


def _is_hard_math_line(text: str) -> bool:
    """强数学信号行：无汉字/全角标点、含 = ∫ ∑ 等信号、且不是英文句子（参考文献等）。"""
    s = text.strip()
    if not s or re.search(r"[一-鿿　-〿＀-￯]", s):
        return False
    if _is_page_number_edge(s) or _looks_like_journal_metadata(s):
        return False
    if not _DEBRIS_SIGNAL_RE.search(s):
        return False
    words = [w.lower() for w in re.findall(r"[A-Za-z]{3,}", s)]
    bad = [w for w in words if w not in _DEBRIS_WORD_WHITELIST]
    if len(bad) > 2:
        return False
    # 表格单元格里的英文短语常带个别希腊字母（如 "A fixed fraction β"），
    # 信号正则会被 β 命中而误判成公式——含非白名单单词时要求运算符/数字佐证
    if bad and not re.search(r"[=≈∝∫∑∏≥≤±⋅·0-9]|\b(?:exp|ln|sin|cos)\b", s):
        return False
    return True


def _is_soft_math_fragment(text: str) -> bool:
    """弱碎片行：短、无汉字/全角标点、仅由数学符号/单字母/数字组成（如 ``|``、``κ = 0``）。"""
    s = text.strip()
    if not s or len(s) > 24 or re.search(r"[一-鿿　-〿＀-￯]", s):
        return False
    if _is_page_number_edge(s) or _looks_like_journal_metadata(s):
        return False
    if len(re.findall(r"[A-Za-z]{3,}", s)) > 1:
        return False
    # 纯英文单词（如表格里的 "irradiance"、"vessel"）字符集能匹配但毫无数学信号，
    # 排除——否则表格文字会被聚类成公式区域
    if re.search(r"[A-Za-z]{3,}", s) and not _DEBRIS_SIGNAL_RE.search(s):
        return False
    return bool(
        re.fullmatch(
            r"[\sA-Za-z0-9+\-\u2212\u2013*/=()(){}\[\]|^_.,:;·⋅×÷<>≤≥≈∝∫∑∏√±∞'′°"
            r"πΠμΜθΘλΛωΩψΨφΦσΣκΚβΒαΑγΓζΖρΡνΝξΞτΤδΔεΕ]+",
            s,
        )
    )


def _insert_math_region_markers(
    page: Any,
    page_text: str,
    regions: list[dict[str, Any]],
    markers: list[str],
) -> tuple[str, set[str]]:
    """按几何阅读顺序把公式区域占位符插到正文对应位置，返回 (新文本, 已插入的标记集合)。

    碎片多重集替换的「首个命中行」锚定在共享碎片行（如 |||、|、z）的页面上
    会把多个公式的标记吸到同一位置；这里改用 dict 行 bbox 的阅读顺序
    （双栏按栏隔离）定位：区域之后的第一条散文行作为锚点，标记插到其前。
    """
    try:
        text_dict = page.get_text("dict")
        page_w = float(page.rect.width)
    except Exception:
        return page_text, set()
    dict_lines: list[dict[str, Any]] = []
    for block in text_dict.get("blocks", []):
        if block.get("type") != 0:
            continue
        for line in block.get("lines", []):
            text = "".join(str(s.get("text", "")) for s in line.get("spans", [])).strip()
            bbox = line.get("bbox")
            if text and bbox:
                dict_lines.append({"text": text, "bbox": [float(v) for v in bbox]})
    if not dict_lines:
        return page_text, set()
    two_col = _is_two_column_layout(dict_lines, page_w)

    def key_of(b: list[float]) -> tuple[float, float]:
        if two_col:
            col = 0.0 if (b[0] + b[2]) / 2 < page_w / 2 else 1.0
            return (col, b[1])
        return (0.0, b[1])

    # 所有区域的碎片行 norm 多重集：锚点必须是散文行，不能是任一公式的零件
    frag_norms: set[str] = set()
    for r in regions:
        for ln in r.get("lines") or []:
            k = _norm_line(ln)
            if k:
                frag_norms.add(k)

    def is_prose(ln: dict[str, Any]) -> bool:
        text = ln["text"]
        norm = _norm_line(text)
        if len(norm) < 8 or norm in frag_norms:
            return False
        if re.search(r"[一-鿿]", text):
            return True
        words = [w.lower() for w in re.findall(r"[A-Za-z]{4,}", text)]
        return any(w not in _DEBRIS_WORD_WHITELIST for w in words)

    order = sorted(range(len(regions)), key=lambda i: key_of(regions[i]["bbox"]))
    lines = page_text.split("\n")
    cursor = 0
    inserts: dict[int, list[str]] = {}  # 锚点行索引 -> 按阅读顺序排列的标记
    for i in order:
        rb = regions[i]["bbox"]
        rcol = key_of([rb[0], rb[3], rb[2], rb[3]])[0]
        rbottom = rb[3]
        best: dict[str, Any] | None = None
        for ln in dict_lines:
            if not is_prose(ln):
                continue
            # 锚点须与区域同栏、且行底伸到区域底部之下：期刊排版行距紧，
            # 紧邻区域的正文行 bbox 常与公式编号行交叠（行顶高于区域底部），
            # 只按行顶比较会把这条正文行漏掉，锚点退到下一行、标记偏一行
            if key_of(ln["bbox"])[0] != rcol or ln["bbox"][3] <= rbottom + 1.0:
                continue
            if best is None or key_of(ln["bbox"]) < key_of(best["bbox"]):
                best = ln
        if best is None:
            continue  # 页尾公式等情形放弃锚定，由调用方按移除位置回退插入
        anchor_norm = _norm_line(best["text"])
        pos = -1
        for j in range(cursor, len(lines)):
            if _norm_line(lines[j]) == anchor_norm:
                pos = j
                break
        if pos < 0:
            for j in range(len(lines)):
                if _norm_line(lines[j]) == anchor_norm:
                    pos = j
                    break
        if pos < 0:
            continue
        inserts.setdefault(pos, []).append(markers[i])
        cursor = pos  # 阅读顺序单调：下一区域的锚点从该位置之后找
    for pos in sorted(inserts, reverse=True):
        lines[pos:pos] = inserts[pos]
    placed = {m for ms in inserts.values() for m in ms}
    return "\n".join(lines), placed


def _remove_math_region_lines(
    page_text: str, region_lines: list[str]
) -> tuple[str, list[str], int | None]:
    """从正文移除公式区域的碎片行（多重集逐行匹配）。

    返回 (新文本, 剥离出的行尾编号, 首个移除行在剩余文本中的行号)。
    碎片行尾粘着的公式编号（N）一并剥离，供调用方挂到公式记录上。
    """
    targets: dict[str, int] = {}
    for ln in region_lines:
        key = _norm_line(ln)
        if key:
            targets[key] = targets.get(key, 0) + 1
    if not targets:
        return page_text, [], None
    out: list[str] = []
    tail_nums: list[str] = []
    first_removed: int | None = None
    for ln in page_text.split("\n"):
        key = _norm_line(ln)
        if key and targets.get(key):
            targets[key] -= 1
            if first_removed is None:
                first_removed = len(out)
            continue
        # 行尾带编号的碎片行（如 ``ext) ⋅z] （5）``）：剥掉编号后命中目标则整行移除
        m = _EQ_NUM_TAIL_RE.search(ln)
        if m:
            head = _norm_line(ln[: m.start()])
            if head and targets.get(head):
                targets[head] -= 1
                tail_nums.append(m.group(1))
                if first_removed is None:
                    first_removed = len(out)
                continue
        out.append(ln)
    return "\n".join(out), tail_nums, first_removed


def _formula_candidates_from_pymupdf(doc: Any) -> list[dict[str, Any]]:
    formulas: list[dict[str, Any]] = []
    for page_index, page in enumerate(doc, start=1):
        text_dict = page.get_text("dict")
        for block in text_dict.get("blocks", []):
            if block.get("type") != 0:
                continue
            for line in block.get("lines", []):
                spans = line.get("spans", [])
                if not spans:
                    continue
                line_text = "".join(span.get("text", "") for span in spans).strip()
                # 利用字体信息辅助公式识别：公式通常使用 italic/oblique 字体
                has_italic = any((span.get("flags", 0) or 0) & 2 for span in spans)
                has_symbol_font = any(
                    "symbol" in str(span.get("font", "")).lower() for span in spans
                )
                candidate = _formula_record(line_text, page_index, "pymupdf-line")
                if candidate:
                    candidate["_font_italic"] = has_italic
                    candidate["_font_symbol"] = has_symbol_font
                    formulas.append(candidate)
                elif has_italic and len(line_text) >= 8:
                    # 斜体行但启发式未识别：尝试宽松匹配
                    candidate = _formula_record(line_text, page_index, "pymupdf-italic")
                    if candidate:
                        formulas.append(candidate)
    # 补充多行公式合并检测（对断裂的公式行进行合并识别）
    lines = [f["raw"] for f in formulas]
    formulas.extend(_multiline_formula_candidates(lines))
    return formulas


def _formula_candidates_from_pdfplumber(content: bytes) -> list[dict[str, Any]]:
    try:
        import pdfplumber
    except Exception:
        return []

    formulas: list[dict[str, Any]] = []
    try:
        with pdfplumber.open(BytesIO(content)) as pdf:
            for page_index, page in enumerate(pdf.pages, start=1):
                text = page.extract_text(x_tolerance=1, y_tolerance=3) or ""
                for line in text.splitlines():
                    candidate = _formula_record(
                        line.strip(), page_index, "pdfplumber-line"
                    )
                    if candidate:
                        formulas.append(candidate)
    except Exception:
        return []
    return formulas


def _formula_candidates_from_text(text: str) -> list[dict[str, Any]]:
    formulas: list[dict[str, Any]] = []
    lines = [line.strip() for line in text.splitlines()]
    for line in lines:
        candidate = _formula_record(line.strip(), None, "pdf-text-line")
        if candidate:
            formulas.append(candidate)
    formulas.extend(_multiline_formula_candidates(lines))
    return formulas


def _multiline_formula_candidates(lines: list[str]) -> list[dict[str, Any]]:
    formulas: list[dict[str, Any]] = []
    index = 0
    while index < len(lines):
        line = lines[index].strip()
        if not _starts_formula_block(line):
            index += 1
            continue
        block = [line]
        cursor = index + 1
        while (
            cursor < len(lines)
            and len(block) < 14
            and _continues_formula_block(lines[cursor])
        ):
            block.append(lines[cursor].strip())
            cursor += 1
        raw = "\n".join(item for item in block if item)
        reconstructed = _reconstruct_multiline_formula(block)
        if reconstructed and _looks_like_formula(reconstructed):
            formulas.append(
                {
                    "text": raw,
                    "raw": raw,
                    "formula": _pdf_formula_text_to_latex(
                        _normalize_pdf_formula_text(reconstructed)
                    ),
                    "source": "pdf-text-block",
                }
            )
            index = cursor
        else:
            index += 1
    return formulas


def _starts_formula_block(line: str) -> bool:
    text = line.strip()
    return bool(
        text
        and re.search(r"[=≈≤≥<>]|\\(?:frac|int|sum)|[∫∑]", text)
        and _math_line_score(text) >= 2
    )


def _continues_formula_block(line: str) -> bool:
    text = line.strip()
    if not text:
        return False
    if len(text) <= 14 and re.fullmatch(r"[A-Za-z0-9+\-*/(){}\[\],.\\_^= ]+", text):
        return True
    return (
        _math_line_score(text) >= 1 and len(re.findall(r"[\u3400-\u9fff]", text)) == 0
    )


def _math_line_score(line: str) -> int:
    return len(
        re.findall(
            r"\\[A-Za-z]+|[=≈≤≥<>+\-*/^_(){}\[\]∫∑√πΠμΜθΘλΛωΩψΨφΦℏħ]|[A-Za-z]_[A-Za-z0-9]|[A-Za-z]\^[A-Za-z0-9]",
            line,
        )
    )


def _reconstruct_multiline_formula(block: list[str]) -> str:
    compact = " ".join(line.strip() for line in block if line.strip())
    compact = re.sub(r"\s+", " ", compact)
    sqrt = "\u221a"
    hbar = "\u210f"
    pi = "\u03c0"
    compact = compact.replace(
        f"(1/ {sqrt} 2{pi}{hbar} )", rf"\frac{{1}}{{\sqrt{{2{pi}{hbar}}}}}"
    )
    compact = compact.replace(
        f"(1 / {sqrt} 2{pi}{hbar} )", rf"\frac{{1}}{{\sqrt{{2{pi}{hbar}}}}}"
    )
    compact = re.sub(
        rf"\(\s*1\s*/\s*{sqrt}\s*([^()]+?)\s*\)", r"\\frac{1}{\\sqrt{\1}}", compact
    )
    compact = re.sub(
        rf"e\s*\(\s*-\s*ip\s*x\s*{hbar}\s*\)", rf"e^{{-ipx/{hbar}}}", compact
    )
    compact = re.sub(
        rf"e\s*\(\s*-\s*ip\s*x\s*/\s*{hbar}\s*\)", rf"e^{{-ipx/{hbar}}}", compact
    )
    compact = re.sub(r"\s*/\s*", "/", compact)
    compact = re.sub(r"\s*([=,+\-])\s*", r"\1", compact)
    compact = re.sub(r"\s+", "", compact)
    return compact


_MATH_OPERATOR_CHARS = r"[=≈≤≥<>±×÷∑∫√∞∂∆Δ∇ħℏπΠμΜθΘλΛωΩψΨφΦα-ωΑ-Ω∝⋅·~]"


def _looks_like_formula(value: str) -> bool:
    text = value.strip()
    if len(text) < 3:
        return False
    if re.search(
        r"<\/?(?:span|sub|sup|math|mrow|mi|mo|mn)\b|class=", text, re.IGNORECASE
    ):
        return False
    # 增加 ∝、⋅ 等符号检测
    math_tokens = re.findall(
        rf"\\[A-Za-z]+|{_MATH_OPERATOR_CHARS}|\^[{{(]?[A-Za-z0-9\\]+|_[{{(]?[A-Za-z0-9\\]+|/[A-Za-z0-9\\({{]",
        text,
    )
    ascii_words = re.findall(r"[A-Za-z]{3,}", text)
    # 增加 ∝、\propto、\exp、\ln、\log 判定
    has_relation = bool(re.search(r"[=≈≤≥<>∝~]|\\(?:leq|geq|approx|propto)", text))
    has_script = bool(re.search(r"[\^_][{{(]?[A-Za-z0-9\\]+", text))
    has_latex = bool(re.search(r"\\[A-Za-z]+", text))
    has_func = bool(re.search(r"\b(?:exp|ln|log|sin|cos|det)\b", text))
    if (
        (has_latex or has_func)
        and len(math_tokens) >= 1
        and (has_relation or has_script)
    ):
        return True
    return (
        len(math_tokens) >= 2
        and (has_relation or has_script or has_latex or has_func)
        and len(ascii_words) <= 6
    )


_FORMULA_TRAILING_NUM_RE = re.compile(
    r"[\(（]\s*(\d{1,3}(?:[a-zA-Z]|\-\d+)?)\s*[\)）]\s*$"
)


def _normalize_pdf_formula_text(value: str) -> str:
    # 转换全角、正比符与点乘（⋅ U+22C5 与 · U+00B7 两种点乘写法）
    value = value.replace("∝", r"\propto ").replace("⋅", r"\cdot ").replace("·", r"\cdot ")
    value = value.translate(
        {
            ord("＝"): "=",
            ord("−"): "-",
            ord("–"): "-",
            ord("≥"): r"\geq",
            ord("≤"): r"\leq",
            ord("≈"): r"\approx",
            ord("∂"): r"\partial",
            ord("Δ"): r"\Delta",
            ord("∇"): r"\nabla",
            ord("∑"): r"\sum",
            ord("∫"): r"\int",
            ord("π"): r"\pi",
            ord("μ"): r"\mu",
            ord("θ"): r"\theta",
            ord("λ"): r"\lambda",
            ord("ω"): r"\omega",
            ord("Ω"): r"\Omega",
            ord("ψ"): r"\psi",
            ord("φ"): r"\phi",
            ord("ℏ"): r"\hbar",
        }
    )
    return value.replace("\u00a0", " ").strip()


def _pdf_formula_text_to_latex(value: str) -> str:
    text = re.sub(r"\s+", " ", value.strip())
    # 剥离行末公式编号（如 （1）、(2)）
    text = _FORMULA_TRAILING_NUM_RE.sub("", text).strip()
    # 规范化常见参数与上下标
    text = re.sub(r"\\mu\s*(t|ext|bs|s|a|fit)\b", r"\\mu_{\1}", text)
    text = re.sub(r"\b([ILF])\s*0\b", r"\1_0", text)
    text = re.sub(r"\b([zh])\s*([0hR]|RD)\b", r"\1_{\2}", text)
    text = re.sub(r"\\Delta\s*(z|\\lambda)\b", r"\\Delta \1", text)
    text = re.sub(r"\bexp\s*\(", r"\\exp(", text)
    text = re.sub(r"\bln\s*\[", r"\\ln[", text)
    text = re.sub(r"\b4\s*\\pi", r"4\\pi", text)
    text = re.sub(r"\b2\s*\\pi", r"2\\pi", text)
    text = re.sub(r"\b([A-Za-z]+)_([A-Za-z0-9]+)\b", r"\1_{\2}", text)
    text = re.sub(r"\b([A-Za-z]+)\^([A-Za-z0-9]+)\b", r"\1^{\2}", text)
    return text


def _formula_record(raw: str, page: int | None, source: str) -> dict[str, Any] | None:
    raw_clean = raw.strip()
    num_match = _FORMULA_TRAILING_NUM_RE.search(raw_clean)
    formula_number = num_match.group(1) if num_match else None

    normalized = _normalize_pdf_formula_text(raw_clean)
    if not _looks_like_formula(normalized):
        return None
    latex = _pdf_formula_text_to_latex(normalized)
    record: dict[str, Any] = {
        "text": normalized,
        "raw": raw_clean,
        "formula": latex,
        "source": source,
    }
    if formula_number:
        record["number"] = formula_number
    if page is not None:
        record["page"] = page
    return record


def _dedupe_formulas(formulas: list[dict[str, Any]]) -> list[dict[str, Any]]:
    deduped: list[dict[str, Any]] = []
    seen: set[str] = set()
    block_texts = [
        re.sub(r"\s+", "", _normalize_pdf_formula_text(str(item.get("text") or "")))
        for item in formulas
        if "\n" in str(item.get("text") or "")
    ]
    for formula in formulas:
        raw_key = re.sub(
            r"\s+", "", _normalize_pdf_formula_text(str(formula.get("text") or ""))
        )
        if "\n" not in str(formula.get("text") or "") and any(
            raw_key and raw_key in block for block in block_texts
        ):
            continue
        key = re.sub(
            r"\s+", "", str(formula.get("formula") or formula.get("text") or "")
        )
        if formula.get("image"):
            # 图片公式每条记录对应独立渲染图，文本相同也不可去重
            # （论文中同一公式多次出现时，第二次的占位符需要自己的记录）
            key = f"{key}@{len(deduped)}"
        if not key or key in seen:
            continue
        seen.add(key)
        deduped.append(formula)
    # 书籍类动辄数百条公式；截图公式排在前面的前提下放宽到 300
    return deduped[:300]


def _docx(filename: str, mime_type: str, content: bytes) -> ExtractedSource:
    from docx import Document

    doc = Document(BytesIO(content))
    text = "\n".join(paragraph.text for paragraph in doc.paragraphs).strip()
    return ExtractedSource(
        filename=filename,
        mime_type=mime_type,
        size_bytes=len(content),
        text=text,
        status="ok" if text else "empty",
        error_message=None if text else "Word 文档未提取到文本。",
    )


def _pptx(filename: str, mime_type: str, content: bytes) -> ExtractedSource:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    deck = Presentation(BytesIO(content))
    parts: list[str] = []
    figures: list[dict[str, Any]] = []

    for slide_index, slide in enumerate(deck.slides, start=1):
        slide_parts: list[str] = [f"[[SLIDE:{slide_index}]]"]
        slide_title = ""

        # 1. 结构化大纲抽取：识别幻灯片标题并标记为二级章节节点
        if slide.shapes.title and slide.shapes.title.text and slide.shapes.title.text.strip():
            slide_title = slide.shapes.title.text.strip()
            slide_parts.append(f"[[H:2]] {slide_title}")

        # 2. 遍历形状（抽取图片、表格结构及正文文本）
        for shape in slide.shapes:
            if slide.shapes.title and shape == slide.shapes.title:
                continue

            # 抽取幻灯片插图
            if getattr(shape, "shape_type", None) == MSO_SHAPE_TYPE.PICTURE or hasattr(shape, "image"):
                try:
                    img = shape.image
                    img_bytes = img.blob
                    img_ext = img.ext.lower() if getattr(img, "ext", None) else "png"
                    mime = "image/jpeg" if img_ext in ("jpg", "jpeg") else f"image/{img_ext}"
                    fig_id = f"fig{len(figures)}"
                    data_uri = f"data:{mime};base64," + base64.b64encode(img_bytes).decode("ascii")
                    caption_text = f"幻灯片 {slide_index}" + (f" · {slide_title}" if slide_title else " 插图")
                    figures.append({
                        "id": fig_id,
                        "page": slide_index - 1,
                        "caption": caption_text,
                        "image": data_uri,
                        "width": getattr(shape, "width", 0),
                        "height": getattr(shape, "height", 0),
                        "source": "pptx",
                    })
                    slide_parts.append(f"[[FIG:{fig_id}]]")
                except Exception:
                    pass

            lines = _pptx_shape_text(shape)
            if lines:
                slide_parts.extend(lines)

        # 3. 抽取演讲者备注
        if getattr(slide, "has_notes_slide", False):
            notes_frame = getattr(slide.notes_slide, "notes_text_frame", None)
            notes = (notes_frame.text or "").strip() if notes_frame is not None else ""
            if notes:
                slide_parts.append(f"备注：{notes}")

        if len(slide_parts) > 1:
            parts.append(f"第 {slide_index} 页\n" + "\n".join(slide_parts))

    text = "\n\n".join(parts).strip()
    return ExtractedSource(
        filename=filename,
        mime_type=mime_type,
        size_bytes=len(content),
        text=text,
        figures=figures,
        status="ok" if text else "empty",
        error_message=None if text else "PPT 未提取到文本。",
    )


def _pptx_shape_text(shape) -> list[str]:
    lines: list[str] = []
    # 格式化表格：生成标准 Markdown 表格，保留行列语义，避免平铺失真
    if getattr(shape, "has_table", False):
        table_rows: list[list[str]] = []
        for row in shape.table.rows:
            cells = [cell.text.replace("\n", " ").strip() for cell in row.cells]
            table_rows.append(cells)
        if table_rows:
            num_cols = max(len(r) for r in table_rows)
            padded = [r + [""] * (num_cols - len(r)) for r in table_rows]
            lines.append("\n| " + " | ".join(padded[0]) + " |")
            lines.append("| " + " | ".join(["---"] * num_cols) + " |")
            for row in padded[1:]:
                lines.append("| " + " | ".join(row) + " |")
            lines.append("")
    if getattr(shape, "has_text_frame", False):
        text = shape.text_frame.text.strip()
        if text:
            lines.append(text)
    if hasattr(shape, "shapes"):
        for child in shape.shapes:
            lines.extend(_pptx_shape_text(child))
    return lines


# ---- PDF 截图兜底：当图片提取失败时，直接渲染页面区域为 PNG ----


def render_page_region(
    content: bytes, page_index: int, bbox: list[float] | None = None, dpi: int = 200
) -> str | None:
    """渲染 PDF 指定页面区域（或整页）为 data URI PNG，供图片提取失败时兜底。

    Args:
        content: PDF 原始字节。
        page_index: 0-based 页码。
        bbox: [x0, y0, x1, y1] 裁剪区域，None 表示整页。
        dpi: 渲染分辨率。

    Returns:
        "data:image/png;base64,..." 或 None（渲染失败）。
    """
    try:
        import pymupdf as fitz
    except Exception:
        import fitz

    try:
        doc = fitz.open(stream=content, filetype="pdf")
    except Exception:
        return None

    try:
        with doc:
            if page_index < 0 or page_index >= len(doc):
                return None
            page = doc[page_index]

            if bbox and len(bbox) == 4:
                tx0, ty0, tx1, ty1 = [float(v) for v in bbox]
                pw = float(page.rect.width)
                ph = float(page.rect.height)
                pad = 8.0
                gap = 1.5  # 公式与最近邻行之间保留的最小安全间距（pt）
                # 逐文本行聚合成 AABB，供上下边/左右边判定最近的邻接行。
                # 公式自身行的中心落在区域内（能被上方/下方顶到 bbox 边缘），
                # 邻接正文/页眉行的中心落在区域外，据此天然区分，避免漏判。
                line_boxes: list[list[float]] = []
                for blk in page.get_text("dict").get("blocks", []):
                    if blk.get("type") != 0:
                        continue
                    for ln in blk.get("lines", []):
                        lx0 = ly0 = 1e9
                        lx1 = ly1 = -1e9
                        for sp in ln.get("spans", []):
                            if not (sp.get("text") or "").strip():
                                continue
                            sb = [float(v) for v in sp["bbox"]]
                            lx0 = min(lx0, sb[0]); ly0 = min(ly0, sb[1])
                            lx1 = max(lx1, sb[2]); ly1 = max(ly1, sb[3])
                        if lx1 >= lx0 and ly1 >= ly0:
                            line_boxes.append([lx0, ly0, lx1, ly1])
                # 每一条边只向留白方向扩展，绝不越过最近的非公式文本行，
                # 从而不把与公式仅隔数 pt 的说明文字、页眉横幅重新截进公式图。
                top_neighbor = max(
                    (l[3] for l in line_boxes if (l[1] + l[3]) * 0.5 < ty0),
                    default=None,
                )
                bot_neighbor = min(
                    (l[1] for l in line_boxes if (l[1] + l[3]) * 0.5 > ty1),
                    default=None,
                )
                left_neighbor = max(
                    (l[2] for l in line_boxes if (l[0] + l[2]) * 0.5 < tx0),
                    default=None,
                )
                right_neighbor = min(
                    (l[0] for l in line_boxes if (l[0] + l[2]) * 0.5 > tx1),
                    default=None,
                )
                if top_neighbor is not None:
                    pad_up = max(0.0, ty0 - top_neighbor - gap)
                else:
                    pad_up = pad
                if bot_neighbor is not None:
                    pad_dn = max(0.0, bot_neighbor - ty1 - gap)
                else:
                    pad_dn = pad
                if left_neighbor is not None:
                    pad_lf = max(0.0, tx0 - left_neighbor - gap)
                else:
                    pad_lf = pad
                if right_neighbor is not None:
                    pad_rt = max(0.0, right_neighbor - tx1 - gap)
                else:
                    pad_rt = pad
                x0 = max(0.0, tx0 - min(pad, pad_lf))
                y0 = max(0.0, ty0 - min(pad, pad_up))
                x1 = min(pw, tx1 + min(pad, pad_rt))
                y1 = min(ph, ty1 + min(pad, pad_dn))
                if x1 - x0 < 1 or y1 - y0 < 1:
                    return None
                clip = fitz.Rect(x0, y0, x1, y1)
                matrix = fitz.Matrix(dpi / 72, dpi / 72)
                try:
                    pix = page.get_pixmap(clip=clip, matrix=matrix, alpha=False)
                except Exception:
                    return None
            else:
                matrix = fitz.Matrix(dpi / 72, dpi / 72)
                try:
                    pix = page.get_pixmap(matrix=matrix, alpha=False)
                except Exception:
                    return None

            img_bytes = pix.tobytes("png")
            if not img_bytes:
                return None
            return "data:image/png;base64," + base64.b64encode(img_bytes).decode(
                "ascii"
            )
    except Exception:
        return None


def find_figure_page_by_caption(
    content: bytes, caption: str
) -> tuple[int, list[float]] | None:
    """在 PDF 各页文本中搜索图表标题，返回 (page_index, page_rect) 或 None。

    用于 MinerU 提取的图表缺少 page/bbox 信息时，定位图表所在页面。
    """
    if not caption or not caption.strip():
        return None

    try:
        import pymupdf as fitz
    except Exception:
        import fitz

    try:
        doc = fitz.open(stream=content, filetype="pdf")
    except Exception:
        return None

    try:
        with doc:
            import re

            normalized = re.sub(r"\s+", "", caption.strip())
            if len(normalized) < 3:
                return None

            for page_index, page in enumerate(doc):
                page_text = page.get_text("text")
                page_normalized = re.sub(r"\s+", "", page_text)
                if normalized in page_normalized:
                    rect = [
                        float(page.rect.x0),
                        float(page.rect.y0),
                        float(page.rect.x1),
                        float(page.rect.y1),
                    ]
                    return (page_index, rect)
            return None
    except Exception:
        return None


class _RapidOCREngine:
    """适配 rapidocr_onnxruntime 1.x 与 rapidocr 3.x 两代 API；推理调用串行化。"""

    def __init__(self) -> None:
        import threading

        self._lock = threading.Lock()
        try:
            from rapidocr_onnxruntime import RapidOCR
        except ImportError:
            from rapidocr import RapidOCR

            self._engine = RapidOCR()
            self._new_api = True
        else:
            self._engine = RapidOCR(use_cls=False, rec_batch_num=8)
            self._new_api = False

    def texts(self, image: Any) -> list[str]:
        with self._lock:
            out = self._engine(image)
        if self._new_api:
            return [str(t) for t in (getattr(out, "txts", None) or [])]
        result, _ = out
        return [str(r[1]) for r in (result or []) if len(r) >= 2]


_RAPIDOCR_ENGINE: _RapidOCREngine | None = None
_RAPIDOCR_FAILED = False


def _get_rapidocr_engine() -> _RapidOCREngine | None:
    """进程级懒加载 RapidOCR 单例（模型加载约 1s，extract_source 可能在线程中调用）。"""
    global _RAPIDOCR_ENGINE, _RAPIDOCR_FAILED
    if _RAPIDOCR_ENGINE is None and not _RAPIDOCR_FAILED:
        try:
            _RAPIDOCR_ENGINE = _RapidOCREngine()
        except ImportError:
            _RAPIDOCR_FAILED = True
    return _RAPIDOCR_ENGINE


def _image(filename: str, mime_type: str, content: bytes) -> ExtractedSource:
    engine = _get_rapidocr_engine()
    if engine is None:
        return ExtractedSource(
            filename=filename,
            mime_type=mime_type,
            size_bytes=len(content),
            text="",
            status="ocr_unavailable",
            error_message="未安装 RapidOCR，图片已接收但暂不能本地 OCR。",
        )

    import numpy as np
    from PIL import Image

    image = np.array(Image.open(BytesIO(content)).convert("RGB"))
    text = "\n".join(engine.texts(image)).strip()
    return ExtractedSource(
        filename=filename,
        mime_type=mime_type,
        size_bytes=len(content),
        text=text,
        status="ok" if text else "empty",
        ocr_used=True,
        error_message=None if text else "图片 OCR 未识别到文本。",
    )
